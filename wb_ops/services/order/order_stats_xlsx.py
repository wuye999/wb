# -*- coding: utf-8 -*-
"""
wb_ops `order-stats` 的 xlsx 渲染（从 order_stats.py 拆出，守住单文件 300 行建议线）

输出结构（单个 xlsx 多 Sheet）：
    `汇总` → 每个店铺一页（按凭证/--shops 顺序）→ `订单明细`
统计列（用户指定，勿改顺序）：
    序号 | 商品中文名 | 订单数 | 退货数 | 退货率 | 完单口径退货率 |
    未完结(已创建进行中) | 涉及店铺 | 供应商代码数 | 库存SKU
百分比一律写**数值 + number_format('0.00%')**（可排序/可透视），仅当分母 ≤0 时写文本 "-"。

样式与 `services/support/ratings.py:_export_excel` 保持一致：
表头 微软雅黑 10 加粗 + DDEBF7 填充 + 居中；数据 9 号；冻结窗格 + 自动筛选。
"""
import os
import re
import time
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from wb_ops import config

HEADERS = ["序号", "商品中文名", "订单数", "退货数", "退货率", "完单口径退货率",
           "未完结(已创建进行中)", "涉及店铺", "供应商代码数", "库存SKU"]
WIDTHS = [6, 34, 9, 9, 11, 16, 20, 46, 14, 18]
CENTER_COLS = (1, 3, 4, 5, 6, 7, 9, 10)
LEFT_COLS = (2, 8)
PCT_COLS = (5, 6)

DETAIL_HEADERS = ["店铺", "店铺ID", "订单ID", "订单号", "下单日期", "下单时间", "更新时间",
                  "状态", "口径归类", "取消类型", "商品中文名", "供应商代码", "库存SKU",
                  "WB商品码", "俄文标题", "尺码", "售价(₽)", "仓库", "目的城市"]
DETAIL_WIDTHS = [14, 8, 36, 30, 11, 26, 26, 10, 10, 10, 30, 26, 16, 13, 34, 10, 10, 30, 16]
KIND_CN = {"return": "退货", "pending": "未完结", "done": "完单", "unknown": "未知状态"}

TITLE_SUMMARY, TITLE_DETAIL, TITLE_FALLBACK = "汇总", "订单明细", "店铺"
ILLEGAL_SHEET_CHARS = set('[]:*?/\\')
MAX_SHEET_LEN = 31
DETAIL_WARN_ROWS = 50000

_HEADER_FONT = Font(name="微软雅黑", size=10, bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_CELL_FONT = Font(name="微软雅黑", size=9)
_TOTAL_FONT = Font(name="微软雅黑", size=9, bold=True)
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_META_FONT = Font(name="微软雅黑", size=9, italic=True, color="808080")


# ---------------- 工具（离线可测） ----------------
def safe_sheet_name(raw: Any, used: Optional[set] = None) -> str:
    """Excel Sheet 名规范化：非法字符→`_`、截断 31 字符、大小写不敏感去重。

    `used` 为已占用集合（小写）；传入则原地登记新名字。`汇总`/`订单明细` 由调用方预占。
    """
    used = used if used is not None else set()
    s = "".join("_" if c in ILLEGAL_SHEET_CHARS else c for c in str(raw or ""))
    s = re.sub(r"\s+", " ", s).strip().strip("'")[:MAX_SHEET_LEN].strip() or TITLE_FALLBACK
    base, i = s, 2
    while s.lower() in used:
        suffix = f"({i})"
        s = base[:MAX_SHEET_LEN - len(suffix)].rstrip() + suffix
        i += 1
    used.add(s.lower())
    return s


def _pct_cell(ws: Worksheet, row: int, col: int, value: Optional[float]) -> None:
    """写百分比单元格：None（分母≤0）→ 文本 "-"；否则 0~1 浮点 + '0.00%'。"""
    c = ws.cell(row=row, column=col)
    c.alignment = Alignment(horizontal="center", vertical="center")
    if value is None:
        c.value = "-"
        return
    c.value = float(value)
    c.number_format = "0.00%"


def _rate(num: int, den: int) -> Optional[float]:
    return (num / den) if den and den > 0 else None


def _style_header(ws: Worksheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def _set_widths(ws: Worksheet, widths: Sequence[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------- 统计 Sheet ----------------
def _write_stat_sheet(ws: Worksheet, rows: Sequence[Dict[str, Any]],
                      total_vc: int = 0, shop_count: int = 0) -> int:
    """写一张统计表（表头 + 数据 + 合计行）；返回数据行数。"""
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get("rank", ""), r.get("cn", ""), r.get("orders", 0), r.get("returns", 0),
                   None, None, r.get("pending", 0), "、".join(r.get("shops") or []),
                   r.get("vc_count", 0), r.get("stock_sku", "")])
        i = ws.max_row
        _pct_cell(ws, i, 5, r.get("rate"))
        _pct_cell(ws, i, 6, r.get("rate_done"))

    n = len(rows)
    for row in range(2, n + 2):
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = _CELL_FONT
            if col in CENTER_COLS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=(col == 8))

    # 合计行（两个率用合计重算，不是行内率求平均）
    orders = sum(int(r.get("orders") or 0) for r in rows)
    returns = sum(int(r.get("returns") or 0) for r in rows)
    pending = sum(int(r.get("pending") or 0) for r in rows)
    tr = n + 2
    ws.append(["", "合计", orders, returns, None, None, pending,
               f"共 {shop_count} 家" if shop_count else "", total_vc, ""])
    _pct_cell(ws, tr, 5, _rate(returns, orders))
    _pct_cell(ws, tr, 6, _rate(returns, orders - pending))
    for col in range(1, len(HEADERS) + 1):
        c = ws.cell(row=tr, column=col)
        c.font = _TOTAL_FONT
        c.fill = _TOTAL_FILL

    _style_header(ws, len(HEADERS))
    _set_widths(ws, WIDTHS)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{n + 1}"
    return n


def _append_meta(ws: Worksheet, lines: Sequence[str]) -> None:
    """数据下方追加元信息块（空一行后逐行灰斜体）。"""
    start = ws.max_row + 2
    for k, text in enumerate(lines):
        c = ws.cell(row=start + k, column=1, value=text)
        c.font = _META_FONT
        c.alignment = Alignment(horizontal="left", vertical="center")


# ---------------- 明细 Sheet ----------------
def _write_detail_sheet(ws: Worksheet, rows: Sequence[Dict[str, Any]], max_rows: int = 0) -> int:
    ws.append(DETAIL_HEADERS)
    shown = rows[:max_rows] if max_rows else rows
    for r in shown:
        ws.append([
            r.get("shop_label", ""), r.get("shop_id", ""), r.get("order_id", ""),
            r.get("order_no", ""), r.get("created_day", ""), r.get("created", ""),
            r.get("updated", ""), r.get("status", ""),
            KIND_CN.get(r.get("kind") or "", r.get("kind") or ""),
            r.get("cancel_type", ""), r.get("cn", ""), r.get("vc", ""), r.get("stock_sku", ""),
            r.get("nm_id", ""), r.get("title", ""), r.get("size", ""),
            r.get("price") if r.get("price") is not None else "", r.get("warehouse", ""),
            r.get("dest", ""),
        ])
    for row in range(2, ws.max_row + 1):
        for col in range(1, len(DETAIL_HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = _CELL_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
    _style_header(ws, len(DETAIL_HEADERS))
    _set_widths(ws, DETAIL_WIDTHS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{ws.max_row}"
    return len(shown)


# ---------------- 落盘 ----------------
def _save_targets(out: str, ts: str) -> List[str]:
    """候选落盘路径：--out 优先；否则中文名 → ASCII 别名（环境按文件名拦截时回退）。"""
    if (out or "").strip():
        return [os.path.abspath(out.strip())]
    return [os.path.join(config.LOG_DIR, f"订单动态统计_{ts}.xlsx"),
            os.path.join(config.LOG_DIR, f"order_stats_{ts}.xlsx")]


def write_xlsx(result: Dict[str, Any], out: str = "", with_detail: bool = True,
               detail_max_rows: int = 0) -> str:
    """把 `order_stats.collect()` 的结果写成 xlsx；返回实际路径（全部失败返回 ""）。"""
    rows = result.get("rows") or []
    ws_meta = [
        f"统计区间：{result.get('begin', '')} ~ {result.get('end', '')}（按订单下单时间 created，闭区间）",
        f"时区：{result.get('timezone', '')}",
        f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}",
        "口径：订单数=明细行数（一行=一件）；退货=status∈{cancel,return,returnDefective}；"
        "未完结=created；退货率=退货数/订单数；完单口径退货率=退货数/(订单数-未完结)",
        "说明：完单口径分母≤0 时显示 '-'；未知状态计入订单数但不归类；"
        "(中文名未知) 桶恒排最后",
    ]
    if result.get("unknown_status"):
        ws_meta.append("未知状态：" + "、".join(f"{k}={v}" for k, v in result["unknown_status"].items()))
    if result.get("failed_shops"):
        ws_meta.append("失败店铺（数据缺失）：" + "、".join(p["label"] for p in result["failed_shops"]))
    if result.get("out_of_window"):
        ws_meta.append(f"接口按 updated 粗筛后，另有 {result['out_of_window']} 行下单日期在窗口外，"
                       f"已按 created 剔除（正常现象）")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = TITLE_SUMMARY
    used = {TITLE_SUMMARY.lower(), TITLE_DETAIL.lower()}
    total_vc = len({r.get("vc") for r in rows if r.get("vc")})
    _write_stat_sheet(ws0, result.get("by_cn") or [], total_vc=total_vc,
                      shop_count=len(result.get("by_shop") or []))
    _append_meta(ws0, ws_meta)

    by_cn_by_shop = result.get("by_cn_by_shop") or {}
    for p in result.get("per_shop") or []:
        if p.get("error") or not by_cn_by_shop.get(p["label"]):
            continue
        ws = wb.create_sheet(safe_sheet_name(p["label"], used))
        _write_stat_sheet(ws, by_cn_by_shop[p["label"]], total_vc=0, shop_count=1)

    detail_n = 0
    if with_detail and rows:
        if len(rows) > DETAIL_WARN_ROWS:
            print(f"  [提示] 订单明细 {len(rows)} 行，文件可能较大（可加 --no-detail 或 --detail-max-rows）")
        ws_d = wb.create_sheet(TITLE_DETAIL)
        detail_n = _write_detail_sheet(ws_d, rows, detail_max_rows)

    ts = time.strftime("%Y%m%d_%H%M%S")
    targets = _save_targets(out, ts)
    last_err = None
    for k, path in enumerate(targets):
        try:
            os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
            wb.save(path)
            if k > 0:
                print(f"  [提示] 首选文件名被环境拦截，已改用 {os.path.basename(path)}")
            return path
        except OSError as ex:      # 含 PermissionError（安全中心按文件名/目录拦截）
            last_err = ex
            continue
    print(f"  [警告] xlsx 保存失败（{type(last_err).__name__}: {last_err}）；如需落盘请用 --out 指定路径")
    return ""
