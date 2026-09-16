# -*- coding: utf-8 -*-
"""
wb_ops 价格映射表 Excel 生成器 (mapping_excel)
封装 8-Sheet 聚合映射表构建、样式渲染、超链接绑定与店铺覆盖矩阵输出。
"""
import glob
import json
import os
import re
from collections import defaultdict
import openpyxl
from wb_ops import config

IMG_COL = 9  # 主图链接列（做超链接）
NMID_COL = 16  # WB商品码列（做超链接）


def stock_summary(r):
    """库存数量：'999'；多规格/多仓库不同数量用逗号分隔（去掉 warehouseId 前缀）"""
    seen, parts = set(), []
    for s in (r.get("sizeList") or []):
        for st in (s.get("stockList") or []):
            am = st.get("amount")
            if am is not None and am not in seen:
                seen.add(am)
                parts.append(str(am))
    return ", ".join(parts)



def detail_headers(main_sid=None):
    """映射总表 16 列（13 主数据 + 创建时间/更新时间/WB商品码） + 店铺覆盖列"""
    sid = main_sid if main_sid is not None else (config.MAIN_SHOP or "")

    return [
        '产品中文名', 'vendorCode', '双倍售价', f'店铺{sid}价格(CNY)', '折扣%', 'club折扣%',
        '库存', '俄文标题', '主图链接', '尺寸长(cm)', '尺寸宽(cm)', '尺寸高(cm)', '毛重(kg)',
        '创建时间', '更新时间', 'WB商品码',
    ]


def row_for(item, vc, c, match="", status="", stock_summary_fn=None):
    """生成一行映射数据（16 列）。item 的售价键兼容 doublePrice（核对JSON）与 dp（补充条目）"""
    dp = item.get("doublePrice", item.get("dp"))
    if c is None:
        return [item.get("cn"), vc, dp, "", "", "", "", "", "", "", "", "", "", "", "", ""]
    stock_val = c.get("_stock")
    if not stock_val and stock_summary_fn:
        stock_val = stock_summary_fn(c)
    return [
        item.get("cn"), vc, dp, c.get("price"),
        c.get("discount"), c.get("clubDiscount"), stock_val or "",
        c.get("title"), c.get("repImg") or "",
        c.get("dimensionsLength"), c.get("dimensionsWidth"), c.get("dimensionsHeight"),
        c.get("dimensionsWeightBrutto"),
        c.get("createAt"), c.get("updateAt"), c.get("nmId"),
    ]


def build_xlsx(result_list, extra_entries, bcs, excluded_vcs=None, shop_coverage=None, shops_meta=None,
               main_sid=None, stock_summary_fn=None, load_boss_fn=None):
    """生成映射表 xlsx（8 Sheet）。"""
    if stock_summary_fn is None:
        from . import mapping
        stock_summary_fn = mapping.stock_summary
    if load_boss_fn is None:
        from . import mapping
        load_boss_fn = mapping.load_boss
    if main_sid is None:
        from . import mapping
        main_sid = mapping.shop_id()

    excluded_vcs = excluded_vcs or set()
    shop_coverage = shop_coverage or {}
    shops_meta = shops_meta or [{"id": main_sid, "name": f"shop{main_sid}"}]
    bcs_by_vc = {c["vc"]: c for c in bcs}
    sid_main = main_sid

    for c in bcs:
        shop_coverage.setdefault(c["vc"], {})[sid_main] = {
            "price": c["price"], "stock": stock_summary_fn(c),
            "nmId": c.get("nmId"), "createAt": c.get("createAt"), "updateAt": c.get("updateAt")}
    picked = set()
    for item in result_list:
        picked.update(item.get("vendorCodes") or [])
    for item in extra_entries:
        picked.update(item.get("vendorCodes") or [])

    vc2boss = defaultdict(list)
    for item in list(result_list) + list(extra_entries):
        for vc in item.get("vendorCodes") or []:
            vc2boss[vc].append(item.get("cn") or "")
    multi = {vc: bs for vc, bs in vc2boss.items() if len(bs) > 1}

    vc_cn = {}
    for item in list(result_list) + list(extra_entries):
        for vc in item.get("vendorCodes") or []:
            vc_cn.setdefault(vc, item.get("cn") or "")

    def rep(c, key, cov):
        v = c.get(key) if c is not None else None
        if v is None and cov:
            d = cov.get(sid_main) or next(iter(cov.values()))
            v = d.get(key)
        return v

    def enrich(vc, c):
        cov = shop_coverage.get(vc, {})
        if c is None:
            if not cov:
                return None
            d0 = next((d for d in cov.values() if d.get("price") is not None), None)
            if d0 is None:
                d0 = next(iter(cov.values()))
            return {"vc": vc, "title": d0.get("title") or "", "price": d0.get("price"),
                    "repImg": d0.get("img") or "", "img": d0.get("img") or "",
                    "_stock": d0.get("stock") or "",
                    "nmId": rep(None, "nmId", cov), "createAt": rep(None, "createAt", cov),
                    "updateAt": rep(None, "updateAt", cov)}
        c = dict(c)
        for sid, d in cov.items():
            if sid == sid_main:
                continue
            if c.get("price") is None and d.get("price") is not None:
                c["price"] = d["price"]
            if not stock_summary_fn(c) and d.get("stock"):
                c["_stock"] = d["stock"]
        for k in ("nmId", "createAt", "updateAt"):
            if c.get(k) is None:
                c[k] = rep(c, k, cov)
        return c

    headers = detail_headers(sid_main)
    wb = openpyxl.Workbook()

    # Sheet1 映射总表
    ws1 = wb.active
    ws1.title = "映射总表"
    ws1.append(headers + ["店铺覆盖"])
    for item in result_list:
        vcs = item.get("vendorCodes") or []
        if not vcs:
            ws1.append(row_for(item, "", None, "", "未匹配", stock_summary_fn=stock_summary_fn) + [""])
            continue
        match = "人工确认" if not item.get("auto") else "自动匹配(参考)"
        for vc in vcs:
            c = enrich(vc, bcs_by_vc.get(vc))
            status = "多重映射冲突" if vc in multi else "已匹配"
            cov = ";".join(str(sid) for sid in sorted(shop_coverage.get(vc, {}).keys())) if vc else ""
            ws1.append(row_for(item, vc, c, match, status, stock_summary_fn=stock_summary_fn) + [cov])
            if c:
                ws1.cell(ws1.max_row, IMG_COL).hyperlink = c.get("repImg") or ""
                if c.get("nmId"):
                    ws1.cell(ws1.max_row, NMID_COL).hyperlink = \
                        f"https://www.wildberries.ru/catalog/{c['nmId']}/detail.aspx?targetUrl=GP"
    for item in extra_entries:
        status = "已补录售价" if item.get("dp") is not None else "待补填售价"
        for vc in item.get("vendorCodes") or []:
            c = enrich(vc, bcs_by_vc.get(vc))
            cov = ";".join(str(sid) for sid in sorted(shop_coverage.get(vc, {}).keys())) if vc else ""
            ws1.append(row_for(item, vc, c, "补充条目(价格带识别)", status, stock_summary_fn=stock_summary_fn) + [cov])
            if c:
                ws1.cell(ws1.max_row, IMG_COL).hyperlink = c.get("repImg") or ""
                if c.get("nmId"):
                    ws1.cell(ws1.max_row, NMID_COL).hyperlink = \
                        f"https://www.wildberries.ru/catalog/{c['nmId']}/detail.aspx?targetUrl=GP"
    ws1.freeze_panes = "A2"
    SIMPLE_WIDTHS = [26, 28, 10, 14, 8, 12, 12, 55, 45, 10, 10, 10, 10, 18, 18, 14, 18]
    for i, w in enumerate(SIMPLE_WIDTHS, 1):
        ws1.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Sheet2 多重映射冲突
    ws2 = wb.create_sheet("多重映射冲突")
    ws2.append(["vendorCode", "被勾选的商品价格表商品数", "涉及的商品价格表商品", "处理建议"])
    if multi:
        for vc, bs in multi.items():
            ws2.append([vc, len(bs), "；".join(bs),
                        "同一vendorCode被多个商品勾选 → 需人工裁决归属（或同款确实共用一个vendorCode时确认保留）"])
    else:
        ws2.append(["", "", "", "无多重映射冲突（本次核对干净）"])
    for i, w in enumerate([26, 16, 60, 50], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Sheet3 未映射商品
    ws3 = wb.create_sheet("未映射商品")
    ws3.append(headers + ["未映射原因"])
    missed = [c for c in bcs if c["vc"] not in picked]
    missed.sort(key=lambda c: c["price"])
    for c in missed:
        if c["vc"] in excluded_vcs:
            reason = "已人工标记为非货盘商品（排除，不参与映射）"
        else:
            reason = "店铺在架但未映射 → 漏配商品候选或非货盘商品，需人工核对"
        row = row_for({"sku": "", "cn": "（未映射）", "doublePrice": None}, c["vc"], c, "", reason, stock_summary_fn=stock_summary_fn)
        ws3.append(row + [reason])
        ws3.cell(ws3.max_row, IMG_COL).hyperlink = c.get("repImg") or ""
    for i, w in enumerate(SIMPLE_WIDTHS + [40], 1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # 已排除清单
    ws_ex = wb.create_sheet("已排除清单")
    ws_ex.append(["vendorCode", "排除原因", "说明"])
    if isinstance(excluded_vcs, dict):
        excl_items = sorted(excluded_vcs.items())
    else:
        excl_items = [(vc, "非货盘商品（人工排除）") for vc in sorted(excluded_vcs)]
    for vc, reason in excl_items:
        ws_ex.append([vc, reason, "不在商品价格表货盘，确认排除；此清单为增量 merge 的排除状态来源"])
    if not excl_items:
        ws_ex.append(["", "", "无已排除商品"])
    for i, w in enumerate([28, 30, 45], 1):
        ws_ex.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Sheet4 待核查清单
    ws4 = wb.create_sheet("待核查清单")
    ws4.append(["商品SKU", "产品中文名", "双倍售价", "问题说明"])
    picked_cn = set()
    for item in list(result_list) + list(extra_entries):
        if item.get("vendorCodes"):
            picked_cn.add(item.get("cn"))
    boss = load_boss_fn()
    for b in boss:
        if b["cn"] not in picked_cn:
            ws4.append([b["sku"], b["cn"], b["dp"],
                        "商品价格表有该商品但映射表无归属 vendorCode → 需人工核对补录（可能未上架/价格特殊/漏配）"])
    for item in extra_entries:
        if item.get("dp") is None:
            note = (item.get("note", "") or "") + " → 需补填双倍售价并确认型号分组"
            ws4.append([item.get("sku"), item.get("cn"), item.get("dp"), note])
    for i, w in enumerate([14, 30, 10, 70], 1):
        ws4.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Sheet5 店铺全量商品
    ws5 = wb.create_sheet("店铺全量商品")
    ws5.append(headers + ["已在映射表"])
    for c in sorted(bcs, key=lambda c: c["price"]):
        row = row_for({"sku": "", "cn": "（店铺商品）", "doublePrice": None}, c["vc"], c, "", "", stock_summary_fn=stock_summary_fn)
        row.append("是" if c["vc"] in picked else "")
        ws5.append(row)
        ws5.cell(ws5.max_row, IMG_COL).hyperlink = c.get("repImg") or ""
    for i, w in enumerate(SIMPLE_WIDTHS + [10], 1):
        ws5.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Sheet6 店铺覆盖矩阵
    ws6 = wb.create_sheet("店铺覆盖矩阵")
    shop_ids = [s["id"] for s in shops_meta]
    ws6.append(["vendorCode", "产品中文名"] + [f"{s['id']}({s['name']})" for s in shops_meta])
    for vc in sorted(picked):
        row = [vc, vc_cn.get(vc, "")]
        for sid in shop_ids:
            d = shop_coverage.get(vc, {}).get(sid)
            row.append(d.get("stock", "") if d else "")
        ws6.append(row)
    for i, w in enumerate([28, 26] + [14] * len(shop_ids), 1):
        ws6.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws6.freeze_panes = "C2"

    # Sheet7 多店价格一致性
    ws7 = wb.create_sheet("多店价格一致性")
    ws7.append(["vendorCode", "产品中文名", f"主店{sid_main}价格", "店铺ID", "该店价格", "说明"])
    n_warn = 0
    for vc in sorted(picked):
        main_price = shop_coverage.get(vc, {}).get(sid_main, {}).get("price")
        if main_price is None:
            continue
        for sid, d in shop_coverage.get(vc, {}).items():
            if sid == sid_main:
                continue
            p = d.get("price")
            if p is not None and abs(float(main_price) - float(p)) > 0.01:
                ws7.append([vc, vc_cn.get(vc, ""), main_price, sid, p,
                            "该店价格与主店不一致 → 人工复核（同品跨店价格应相同）"])
                n_warn += 1
    if n_warn == 0:
        ws7.append(["", "", "", "", "", "无价格不一致（全部店铺同品同价 ✓）"])
    for i, w in enumerate([28, 26, 14, 10, 14, 50], 1):
        ws7.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(config.MAPPING_XLSX)
    n_unmapped = sum(1 for c in bcs if c["vc"] not in picked)
    return len(result_list), len(picked), n_unmapped, len(multi)


SINGLE_SHOP_HEADERS = [
    '产品中文名', 'vendorCode', '双倍售价', '本店售价(CNY)', '折扣%', 'club折扣%',
    '本店库存', '本店nmId', '俄文标题', '主图链接',
    '尺寸长(cm)', '尺寸宽(cm)', '尺寸高(cm)', '毛重(kg)',
    '创建时间', '更新时间', '匹配类型'
]
SINGLE_SHOP_WIDTHS = [20, 24, 10, 12, 8, 10, 10, 14, 32, 14, 10, 10, 10, 10, 18, 18, 12]

SINGLE_UNMAPPED_HEADERS = [
    'vendorCode', '俄文标题', '本店售价(CNY)', '折扣%', '本店库存', '本店nmId', '主图链接', '创建时间'
]
SINGLE_UNMAPPED_WIDTHS = [24, 32, 12, 8, 10, 14, 14, 18]


def build_single_shop_xlsx(sid, shop_name=None, rows=None, registry=None, boss=None, prefix_map=None, out_path=None):
    """构建单店铺映射表（包含该店真实在架商品、实际售价、实际库存、实际nmId等）。
    包含 2 个 Sheet：
      1. 在架映射明细（已归属中文名的商品）
      2. 未映射在架商品（在架但未匹配到中文名的商品）
    """
    if rows is None:
        p_json = config.shop_json_path(sid)
        if not os.path.exists(p_json):
            print(f"  [警告] 店铺 {sid} 无商品快照：{p_json}，跳过构建单店表")
            return 0, 0
        from wb_ops.framework.safe_io import safe_load_json
        d = safe_load_json(p_json, default={})
        rows = d.get("rows", [])

    if registry is None:
        from wb_ops.storage.mapping_repo import load_vc_known, load_vc_override, load_vc_excluded
        known = load_vc_known()
        overrides = load_vc_override()
        excluded = load_vc_excluded()
    else:
        known, overrides, excluded = registry

    if boss is None:
        from wb_ops.storage.mapping_repo import load_boss
        boss = load_boss()
    boss_by_cn = {b["cn"]: b for b in boss if b.get("cn")}

    if prefix_map is None:
        from wb_ops.storage.mapping_repo import load_prefix_map
        prefix_map = load_prefix_map()

    out_file = out_path or config.shop_mapping_xlsx(sid, shop_name)
    os.makedirs(os.path.dirname(out_file), exist_ok=True)

    alive_rows = [r for r in rows if not r.get("trashedAt")]

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "在架映射明细"
    ws1.append(SINGLE_SHOP_HEADERS)

    ws2 = wb.create_sheet("未映射在架商品")
    ws2.append(SINGLE_UNMAPPED_HEADERS)

    n_mapped, n_unmapped = 0, 0
    known_updated = False
    for r in alive_rows:
        vc = r.get("vendorCode")
        if not vc:
            continue
        sl = r.get("sizeList") or []
        price = sl[0].get("price") if sl else None
        disc = r.get("discount")
        club = r.get("clubDiscount")
        stock = stock_summary(r)
        nm_id = r.get("nmId")
        title = r.get("title") or ""
        img = r.get("repImg") or ""
        L = r.get("dimensionsLength")
        W = r.get("dimensionsWidth")
        H = r.get("dimensionsHeight")
        weight = r.get("dimensionsWeightBrutto")
        create_at = r.get("createAt")
        update_at = r.get("updateAt")

        cn = ""
        dp = None
        match_type = ""

        # 匹配归属优先级：overrides > known > prefix_map
        if vc in overrides:
            cn = overrides[vc].get("cn") or ""
            dp = overrides[vc].get("dp")
            match_type = "人工纠偏"
        elif vc in known:
            cn = known[vc].get("cn") or ""
            dp = known[vc].get("dp")
            match_type = "已归属"
            if L is None and known[vc].get("L"):
                L, W, H, weight = known[vc].get("L"), known[vc].get("W"), known[vc].get("H"), known[vc].get("weight")
        else:
            m = re.match(config.VC_PREFIX_RE, vc or "")
            if m:
                pfx = m.group(1)
                boss_row = prefix_map.get(pfx)
                if boss_row:
                    cn = boss_row.get("cn") or ""
                    dp = boss_row.get("dp")
                    match_type = "前缀自动"
                    known[vc] = {"cn": cn, "dp": dp, "sku": boss_row.get("sku", ""), "source": "前缀自动"}
                    known_updated = True

        # 用 boss 最新售价覆盖双倍售价
        if cn in boss_by_cn:
            b = boss_by_cn[cn]
            if b.get("dp") is not None:
                dp = b["dp"]

        if cn:
            ws1.append([
                cn, vc, dp, price, disc, club, stock, nm_id, title, img,
                L, W, H, weight, create_at, update_at, match_type
            ])
            cur_row = ws1.max_row
            if img:
                ws1.cell(cur_row, 10).hyperlink = img
            if nm_id:
                ws1.cell(cur_row, 8).hyperlink = f"https://www.wildberries.ru/catalog/{nm_id}/detail.aspx?targetUrl=GP"
            n_mapped += 1
        else:
            ws2.append([vc, title, price, disc, stock, nm_id, img, create_at])
            cur_row = ws2.max_row
            if img:
                ws2.cell(cur_row, 7).hyperlink = img
            if nm_id:
                ws2.cell(cur_row, 6).hyperlink = f"https://www.wildberries.ru/catalog/{nm_id}/detail.aspx?targetUrl=GP"
            n_unmapped += 1

    for i, w in enumerate(SINGLE_SHOP_WIDTHS, 1):
        ws1.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws1.freeze_panes = "C2"

    for i, w in enumerate(SINGLE_UNMAPPED_WIDTHS, 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2.freeze_panes = "B2"

    wb.save(out_file)
    wb.close()

    if known_updated:
        from wb_ops.storage.mapping_repo import save_vc_known
        save_vc_known(known)

    return n_mapped, n_unmapped


def load_single_shop_rows(sid, shop_name=None):
    """读取某店铺独立映射表 Excel，返回 {vc: row_dict} 与 [unmapped_vc]"""
    p = config.shop_mapping_xlsx(sid, shop_name)
    if not os.path.exists(p):
        return {}, []
    wb = openpyxl.load_workbook(p, data_only=True)
    mapped = {}
    if "在架映射明细" in wb.sheetnames:
        ws = wb["在架映射明细"]
        headers = [str(c.value or "") for c in ws[1]]
        idx = {h: i for i, h in enumerate(headers)}
        for r in ws.iter_rows(min_row=2, values_only=True):
            vc = r[idx.get("vendorCode", 1)]
            if not vc:
                continue
            mapped[str(vc)] = {
                "cn": r[idx.get("产品中文名", 0)] or "",
                "vc": str(vc),
                "dp": r[idx.get("双倍售价", 2)] if "双倍售价" in idx else None,
                "price": r[idx.get("本店售价(CNY)", 3)] if "本店售价(CNY)" in idx else None,
                "discount": r[idx.get("折扣%", 4)] if "折扣%" in idx else None,
                "club": r[idx.get("club折扣%", 5)] if "club折扣%" in idx else None,
                "stock": r[idx.get("本店库存", 6)] if "本店库存" in idx else None,
                "nmId": r[idx.get("本店nmId", 7)] if "本店nmId" in idx else None,
                "title": r[idx.get("俄文标题", 8)] if "俄文标题" in idx else "",
                "img": r[idx.get("主图链接", 9)] if "主图链接" in idx else "",
                "L": r[idx.get("尺寸长(cm)", 10)] if "尺寸长(cm)" in idx else None,
                "W": r[idx.get("尺寸宽(cm)", 11)] if "尺寸宽(cm)" in idx else None,
                "H": r[idx.get("尺寸高(cm)", 12)] if "尺寸高(cm)" in idx else None,
                "weight": r[idx.get("毛重(kg)", 13)] if "毛重(kg)" in idx else None,
                "createAt": r[idx.get("创建时间", 14)] if "创建时间" in idx else None,
                "updateAt": r[idx.get("更新时间", 15)] if "更新时间" in idx else None,
                "match_type": r[idx.get("匹配类型", 16)] if "匹配类型" in idx else "",
            }
    unmapped = []
    if "未映射在架商品" in wb.sheetnames:
        ws2 = wb["未映射在架商品"]
        for r in ws2.iter_rows(min_row=2, values_only=True):
            if r[0]:
                unmapped.append(str(r[0]))
    wb.close()
    return mapped, unmapped


def list_active_shop_mappings():
    """扫描 data/shops/ 下所有活跃店铺映射表（排除 _archive/ 目录与临时文件）。
    返回 [{'id': sid, 'name': name, 'path': filepath}, ...]"""
    os.makedirs(config.SHOPS_DIR, exist_ok=True)
    files = sorted(glob.glob(os.path.join(config.SHOPS_DIR, "shop_*.xlsx")))
    active = []
    for f in files:
        base = os.path.basename(f)
        if base.startswith("~$") or "_archive" in f:
            continue
        m = re.match(r"^shop_(\d+)(?:_(.+))?\.xlsx$", base)
        if m:
            sid = int(m.group(1))
            name = m.group(2) or f"shop{sid}"
            active.append({"id": sid, "name": name, "path": f})
    return active

