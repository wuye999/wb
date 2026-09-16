# -*- coding: utf-8 -*-
"""
wb_ops 他人映射表解析与包装规格解析 (foreign_table)
处理他人映射表「映射总表」Sheet 的加载、去重、供应商代码校验及包装尺寸/重量兜底推断。
"""
import math
import openpyxl

from wb_ops import common
from wb_ops.services.replicate import replicate

# 他人表「映射总表」Sheet 用到的列名（按表头定位，缺列容错）
COL_CN = "产品中文名"
COL_VC = "vendorCode"
COL_DP = "双倍售价"
COL_TITLE = "俄文标题"
COL_IMG = "主图链接"
COL_L = "尺寸长(cm)"
COL_W = "尺寸宽(cm)"
COL_H = "尺寸高(cm)"
COL_WT = "毛重(kg)"
COL_NM = "WB商品码"  # 他人店铺当前可靠 nmId（抓取上品数据必须项，无此列跳过）

OZON_BADGE = "ozon-card"


def nm_of(vc):
    """从 vendorCode 提取原始 WB 商品码（用于跨账号跨店铺查重与差集比对）"""
    return common.extract_wb_nm(vc)


def vc_tail_badge(vc):
    """提取他人 vc 除去前缀与末段 nm 后的中间修饰段（如 ozon-card 或 WRLINWI/ 等）"""
    import re
    s = str(vc or "").strip()
    if "/" in s:
        m = re.match(r"^BCS-[A-Z]{4}-([^/]+/)\d+$", s)
        if m:
            return m.group(1)
    elif f"-{OZON_BADGE}-" in s:
        return f"{OZON_BADGE}-"
    return ""


def load_foreign(xlsx_path):
    """解析他人映射表「映射总表」Sheet → (items, bad_vcs, no_nm_count)
    items: [{cn, vc, k, nm, dp, title_ru, img, L, W, H, weight}]，同原始 WB 码 k 去重（优先双倍售价非空行）
      k: = common.extract_wb_nm(vc)，供应商代码里的原始 WB 商品码（跨店铺防重复上架判断的核心键）
      nm: = 他人表「WB商品码」列（他人店铺当前可靠 nmId，新上架 API 提交抓取的 sku，必须非空且纯数字）
    bad_vcs: 供应商代码无法提取原始 WB 码的异常行
    no_nm_count: 缺少有效 WB商品码（无法用于新接口抓取上品）被跳过的行数
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "映射总表" not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"{xlsx_path} 缺少「映射总表」Sheet（需同项目生成的映射表格式）")
    ws = wb["映射总表"]
    headers = [str(c.value or "") for c in ws[1]]

    def idx(name):
        return headers.index(name) if name in headers else None

    i_cn, i_vc, i_dp = idx(COL_CN), idx(COL_VC), idx(COL_DP)
    i_title, i_img = idx(COL_TITLE), idx(COL_IMG)
    i_l, i_w, i_h, i_wt = idx(COL_L), idx(COL_W), idx(COL_H), idx(COL_WT)
    i_nm = idx(COL_NM)

    if i_vc is None or i_cn is None:
        wb.close()
        raise RuntimeError(f"表头缺少 {COL_VC}/{COL_CN} 列，无法解析")
    if i_nm is None:
        wb.close()
        print(f"[错误] 他人表缺少「{COL_NM}」列，无法读取抓取上品所需的可靠 WB商品码！按规则坚决跳过，无法导入。")
        return [], [], 0

    by_k, bad_vcs = {}, []
    no_nm_count = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        vc = str(r[i_vc] or "").strip()
        cn = str(r[i_cn] or "").strip() if i_cn is not None else ""
        if not vc and not cn:
            continue

        # 1. 提取供应商代码里的原始 WB 商品码（查重防重核心）
        k = nm_of(vc)
        if not k:
            bad_vcs.append(vc)
            continue

        # 2. 提取他人店铺在售 WB商品码（API 抓取上品目标 sku，做严格非空与纯数字校验）
        raw_nm = r[i_nm] if i_nm is not None else None
        nm = str(raw_nm).strip() if raw_nm not in (None, "") else ""
        if not nm or not nm.isdigit():
            no_nm_count += 1
            continue

        item = {
            "cn": cn,
            "vc": vc,
            "k": k,
            "nm": nm,
            "dp": r[i_dp] if i_dp is not None else None,
            "title_ru": str(r[i_title] or "").strip() if i_title is not None else "",
            "img": str(r[i_img] or "").strip() if i_img is not None else "",
            "L": r[i_l] if i_l is not None else None,
            "W": r[i_w] if i_w is not None else None,
            "H": r[i_h] if i_h is not None else None,
            "weight": r[i_wt] if i_wt is not None else None,
        }
        old = by_k.get(k)
        if old is None or (old["dp"] is None and item["dp"] is not None):
            by_k[k] = item

    wb.close()
    return list(by_k.values()), bad_vcs, no_nm_count


def resolve_import_package(item, card_cache=None):
    """
    获取导入上架的包装尺寸与重量：
    1. 优先读取他人表 (item) 的 L, W, H, weight
    2. 若缺失或 <= 0，查商品价格表 (boss_pkg_map) 兜底
    3. 若仍缺失，查 card.json 兜底
    返回: ((L, W, H, weight), None) 或 (None, err_msg)
    """
    length = item.get("L")
    width = item.get("W")
    height = item.get("H")
    weight = item.get("weight")
    cn = item.get("cn") or ""
    nm = item.get("nm")

    def _valid(val):
        try:
            return float(val) > 0
        except (TypeError, ValueError):
            return False

    # 2. 商品价格表兜底
    if not (_valid(length) and _valid(width) and _valid(height) and _valid(weight)):
        bp = replicate.boss_pkg_map().get(cn)
        if bp:
            bl, bw, bh, bwgt = bp
            if not _valid(length):
                length = bl
            if not _valid(width):
                width = bw
            if not _valid(height):
                height = bh
            if not _valid(weight):
                weight = bwgt

    # 3. card.json 兜底
    if not (_valid(length) and _valid(width) and _valid(height) and _valid(weight)):
        card_info = None
        if card_cache is not None and nm in card_cache:
            card_info = card_cache[nm]
        else:
            card_info = replicate.fetch_card_json(nm)
            if card_cache is not None:
                card_cache[nm] = card_info
        if card_info:
            pkg = replicate.parse_package_info(card_info)
            if not _valid(length):
                length = pkg.get("length")
            if not _valid(width):
                width = pkg.get("width")
            if not _valid(height):
                height = pkg.get("height")
            if not _valid(weight):
                weight = pkg.get("weight")

    if not (_valid(length) and _valid(width) and _valid(height) and _valid(weight)):
        return None, f"包装数据缺失（L={length} W={width} H={height} 重={weight}）"

    return (
        math.ceil(float(length)),
        math.ceil(float(width)),
        math.ceil(float(height)),
        round(float(weight), 3)
    ), None
