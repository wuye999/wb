# -*- coding: utf-8 -*-
"""
wb_ops 映射与品名归属仓储 (MappingRepository)
统一管理全局 VC 归属池、纠偏改名、前缀映射及排除清单的安全持久化与内存缓存。
"""
import os
import re
from typing import Dict, List, Any, Tuple, Callable
from .. import config
from ..framework.safe_io import atomic_dump_json, safe_load_json
from .boss_price import (CachedDpReader, floor_price, resolve_formula,
                         to_number, warn_unresolved, boss_columns)
from ..framework.xlsx_headers import (COL_ACTUAL_PRICE, COL_CN, COL_H, COL_L, COL_NM,
                                      COL_SHOP_CNY_RE, COL_VC, COL_W, COL_WT,
                                      HeaderError, find_col, find_col_re,
                                      warn_missing)



class MappingRepository:
    """映射数据仓储"""

    @staticmethod
    def load_vc_known() -> Dict[str, Any]:
        """加载全局已知 VC 归属池 (vc_known.json)"""
        return safe_load_json(config.VC_KNOWN_JSON, default={})

    @staticmethod
    def save_vc_known(data: Dict[str, Any]):
        """原子保存全局已知 VC 归属池"""
        atomic_dump_json(config.VC_KNOWN_JSON, data, indent=2, use_lock=True)

    @staticmethod
    def load_vc_override() -> Dict[str, Any]:
        """加载全局 VC 纠偏改名清单 (vc_override.json)"""
        return safe_load_json(config.VC_OVERRIDE_JSON, default={})

    @staticmethod
    def save_vc_override(data: Dict[str, Any]):
        """原子保存全局 VC 纠偏改名清单"""
        atomic_dump_json(config.VC_OVERRIDE_JSON, data, indent=2, use_lock=True)

    @staticmethod
    def load_vc_excluded() -> Dict[str, Any]:

        """加载全局排除清单 (vc_excluded.json)"""
        return safe_load_json(config.VC_EXCLUDED_JSON, default={})

    @staticmethod
    def save_vc_excluded(data: Dict[str, Any]):
        """原子保存全局排除清单"""
        atomic_dump_json(config.VC_EXCLUDED_JSON, data, indent=2, use_lock=True)


    @staticmethod
    def _resolve_dp(formula, ws, row_idx=None):
        """解析 D 列「实际售价」公式（`=F57*2` / `=F57` / `=$F$57*2` / `=ROUND(F57*1.3,2)`）。

        无法解析返回 None，由 `load_boss()` 再回退 Excel 缓存值（data_only）。
        实现见 `storage/boss_price.py`（受限求值，不执行任意函数）。
        """
        return resolve_formula(formula, ws)

    @classmethod
    def load_boss(cls) -> List[Dict[str, Any]]:
        """商品价格表 → [{idx, sku, cn, dp(实际售价), img, floor, prefix, stock_sku}]

        列全部**按表头名定位**（不依赖列位，用户可随意插列/挪列）：
        产品图片 / 卖家SKU / 产品中文名 / 实际售价 / 尺寸 / 最低售价 / vendorCode前缀码 / 库存SKU。
        「实际售价」可写数值或公式（`=F2*2` / `=F2` / `=$F$2*2` / `=ROUND(F2*1.3,2)`）；
        公式不可解析时回退 Excel 缓存值，仍取不到则 dp/floor=None 并打印 [警告]（不再静默空价）。
        缺「实际售价」列 → HeaderError（fail-fast，避免改价/上架全部跳过）。
        `floor` = math.floor(实际售价)，即店铺整数价口径。"""
        if not os.path.exists(config.BOSS_XLSX):
            return []
        import openpyxl
        wb = openpyxl.load_workbook(config.BOSS_XLSX, data_only=False)
        try:
            ws = wb["Sheet1"]
            cols = boss_columns(ws, config.BOSS_XLSX)
            i_img, i_sku, i_cn = cols["img"], cols["sku"], cols["cn"]
            i_dp, i_prefix, i_stock = cols["dp"], cols["prefix"], cols["stock_sku"]
            cache = CachedDpReader(config.BOSS_XLSX, col=i_dp + 1)
            items: List[Dict[str, Any]] = []
            by_cache: List[Tuple] = []
            missing: List[Tuple] = []

            def val(row, i):
                return row[i].value if i is not None and len(row) > i else None

            try:
                for row in ws.iter_rows(min_row=2):
                    sku = str(val(row, i_sku) or "").strip()
                    cn = str(val(row, i_cn) or "").strip()
                    if not sku and not cn:
                        continue
                    raw = val(row, i_dp)
                    dp = to_number(raw)
                    if dp is None and isinstance(raw, str) and raw.strip().startswith("="):
                        dp = cls._resolve_dp(raw, ws, row[0].row)
                        if dp is None:
                            dp = to_number(cache.get(row[0].row))
                            (by_cache if dp is not None else missing).append(
                                (row[0].row, sku, raw))
                    elif dp is None and raw not in (None, ""):
                        missing.append((row[0].row, sku, raw))
                    items.append({
                        "idx": len(items) + 1,
                        "sku": sku,
                        "cn": cn,
                        "dp": dp,
                        "img": str(val(row, i_img) or "").strip(),
                        "floor": floor_price(dp),
                        "prefix": str(val(row, i_prefix) or "").strip().upper(),
                        "stock_sku": str(val(row, i_stock) or "").strip(),
                    })
            finally:
                cache.close()
        finally:
            wb.close()
        warn_unresolved(config.BOSS_XLSX, by_cache, missing)
        return items

    @staticmethod
    def load_prefix_map() -> Dict[str, Any]:
        """从商品价格表加载 4 位前缀码映射（列按表头名定位）。"""

        if not os.path.exists(config.BOSS_XLSX):
            return {}
        wb = None
        try:
            import openpyxl
            wb = openpyxl.load_workbook(config.BOSS_XLSX, data_only=True)
            ws = wb["Sheet1"]
            cols = boss_columns(ws, config.BOSS_XLSX)
            i_sku, i_cn, i_dp, i_prefix = cols["sku"], cols["cn"], cols["dp"], cols["prefix"]
            pmap = {}
            for row in ws.iter_rows(min_row=2):
                get = lambda i: row[i].value if i is not None and len(row) > i else None
                sku = str(get(i_sku) or "").strip()
                cn = str(get(i_cn) or "").strip()
                if not sku and not cn:
                    continue
                prefix = str(get(i_prefix) or "").strip().upper()
                if not prefix:
                    continue
                dp = get(i_dp)
                if prefix not in pmap:
                    pmap[prefix] = {"sku": sku, "cn": cn, "dp": dp}
            return pmap
        except HeaderError:
            raise                       # 表头错误必须冒泡，不能被下面的宽 except 吞掉
        except Exception:
            return {}
        finally:
            if wb is not None:
                wb.close()

    @classmethod
    def load_boss_details(cls) -> Dict[str, Dict[str, Any]]:
        """从商品价格表加载完整商品详情索引（按 4 位前缀码与中文名双向索引）。

        返回:
            {
                "by_prefix": {prefix: {"sku", "cn", "dp", "dim", "price", "prefix", "stock_sku"}},
                "by_cn": {cn: {"sku", "cn", "dp", "dim", "price", "prefix", "stock_sku"}}
            }
        """
        if not os.path.exists(config.BOSS_XLSX):
            return {"by_prefix": {}, "by_cn": {}}
        wb = None
        try:
            import openpyxl
            wb = openpyxl.load_workbook(config.BOSS_XLSX, data_only=True)
            ws = wb["Sheet1"]
            cols = boss_columns(ws, config.BOSS_XLSX)
            i_sku, i_cn, i_dp = cols["sku"], cols["cn"], cols["dp"]
            i_dim, i_price = cols["dim"], cols["price"]
            i_prefix, i_stock = cols["prefix"], cols["stock_sku"]
            by_prefix = {}
            by_cn = {}
            for row in ws.iter_rows(min_row=2):
                get = lambda i: row[i].value if i is not None and len(row) > i else None
                sku = str(get(i_sku) or "").strip()
                cn = str(get(i_cn) or "").strip()
                if not sku and not cn:
                    continue
                dp = get(i_dp)
                dim = str(get(i_dim) or "").strip()
                price = get(i_price)
                prefix = str(get(i_prefix) or "").strip().upper()
                stock_sku = str(get(i_stock) or "").strip()
                item = {
                    "sku": sku,
                    "cn": cn,
                    "dp": dp,               # 实际售价（价格表「实际售价」列，用户维护；店铺价 = floor）
                    "dim": dim,             # 尺寸
                    "price": price,         # 最低售价
                    "prefix": prefix,       # vendorCode前缀码
                    "stock_sku": stock_sku,
                }
                if prefix and prefix not in by_prefix:
                    by_prefix[prefix] = item
                if cn and cn not in by_cn:
                    by_cn[cn] = item
            return {"by_prefix": by_prefix, "by_cn": by_cn}
        except HeaderError:
            raise                       # 表头错误必须冒泡
        except Exception:
            return {"by_prefix": {}, "by_cn": {}}
        finally:
            if wb is not None:
                wb.close()


    @classmethod
    def build_vc_resolver(cls) -> Tuple[Callable[[str, str], str], Dict[str, str]]:
        """构建毫秒级商品中文名反查解析器。
        
        解析优先级:
        1. vc_override.json (最高优先级纠偏)
        2. vc_known.json (历史已知映射)
        3. 4位前缀码映射 (prefix_map 兜底)
        
        返回: (resolve_cn 函数, 全局已加载的 vc_cn 字典)
        """
        vc_cn: Dict[str, str] = {}

        # 1. vc_known.json
        known = cls.load_vc_known()
        for vc, info in known.items():
            if isinstance(info, dict) and info.get("cn"):
                vc_cn[vc.upper()] = info["cn"]

        # 2. vc_override.json
        override = cls.load_vc_override()
        for vc, info in override.items():
            if isinstance(info, dict) and info.get("cn"):
                vc_cn[vc.upper()] = info["cn"]

        # 3. 前缀映射兜底
        prefix_map = cls.load_prefix_map()
        prefix_re = re.compile(config.VC_PREFIX_RE)

        def resolve_cn(vc: str, default_title: str = "") -> str:
            if not vc:
                return default_title
            v_upper = vc.strip().upper()
            if v_upper in vc_cn:
                return vc_cn[v_upper]
            m = prefix_re.match(v_upper)
            if m:
                pref = m.group(1).upper()
                if pref in prefix_map and prefix_map[pref].get("cn"):
                    return prefix_map[pref]["cn"]
            return default_title

        return resolve_cn, vc_cn

    @classmethod
    def load_mapping_state(cls) -> Tuple[Dict[str, Any], Dict[str, str]]:
        """读现有映射表 xlsx → 增量状态（归属关系 + 已排除清单）。
        返回 (state, excluded)：state={vc:{cn,dp,shop_price,discount,nmId,...}}；excluded={vc:原因}。
        映射表不存在 → 空。自动应用 vc_override.json 纠偏记录（优先级最高）。"""
        state, excluded = {}, {}
        if not os.path.exists(config.MAPPING_XLSX):
            return state, excluded
        import sys
        import time
        retries = 3
        for attempt in range(retries):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(config.MAPPING_XLSX, data_only=True)
                if "映射总表" in wb.sheetnames:
                    ws = wb["映射总表"]
                    headers = [str(c.value or "").strip() for c in ws[1]]
                    if COL_VC in headers:
                        i_cn, i_vc = headers.index(COL_CN), headers.index(COL_VC)
                        # 价格列只认新名「实际售价」；旧表缺列时告警降级（不能 raise：
                        # merge→calc_new_vcs→load_known_vcs 会经过这里，硬失败会阻塞自愈）
                        i_dp = find_col(headers, COL_ACTUAL_PRICE)
                        if i_dp is None:
                            warn_missing("价格映射表·映射总表", config.MAPPING_XLSX, headers,
                                         COL_ACTUAL_PRICE,
                                         hint="dp 全部为空；请运行 wb.py merge 重建映射表")
                        i_sp = find_col_re(headers, COL_SHOP_CNY_RE)
                        i_disc = find_col(headers, "折扣%")
                        i_nm = find_col(headers, COL_NM)
                        i_l = find_col(headers, COL_L)
                        i_w = find_col(headers, COL_W)
                        i_h = find_col(headers, COL_H)
                        i_wt = find_col(headers, COL_WT)
                        for r in ws.iter_rows(min_row=2, values_only=True):
                            vc = r[i_vc]
                            if vc:
                                state[vc] = {
                                    "cn": r[i_cn] or "",
                                    "dp": r[i_dp] if i_dp is not None else None,
                                    "shop_price": r[i_sp] if i_sp is not None else None,
                                    "discount": r[i_disc] if i_disc is not None else None,
                                    "nmId": (r[i_nm] if i_nm is not None else None),
                                    "L": (r[i_l] if i_l is not None else None),
                                    "W": (r[i_w] if i_w is not None else None),
                                    "H": (r[i_h] if i_h is not None else None),
                                    "weight": (r[i_wt] if i_wt is not None else None),
                                }
                if "已排除清单" in wb.sheetnames:
                    ws2 = wb["已排除清单"]
                    for r in ws2.iter_rows(min_row=2, values_only=True):
                        if r[0]:
                            excluded[r[0]] = r[1] or "非货盘商品（人工排除）"
                wb.close()
                break
            except (PermissionError, OSError) as e:
                if attempt == retries - 1:
                    print(f"[警告] 读取映射表 {config.MAPPING_XLSX} 失败（文件可能被占用）: {e}", file=sys.stderr)
                time.sleep(0.1 * (attempt + 1))
            except Exception as e:
                print(f"[警告] 读取映射表 {config.MAPPING_XLSX} 发生异常: {e}", file=sys.stderr)
                break

        # 自动应用全局纠偏更正（高优先级）
        overrides = cls.load_vc_override()
        for vc, ov in overrides.items():
            if vc in state and isinstance(ov, dict) and ov.get("cn"):
                state[vc]["cn"] = ov["cn"]

        return state, excluded


# 模块级便捷访问导出
load_boss = MappingRepository.load_boss
load_prefix_map = MappingRepository.load_prefix_map
load_vc_known = MappingRepository.load_vc_known
save_vc_known = MappingRepository.save_vc_known
load_vc_override = MappingRepository.load_vc_override
save_vc_override = MappingRepository.save_vc_override
load_vc_excluded = MappingRepository.load_vc_excluded
save_vc_excluded = MappingRepository.save_vc_excluded
load_mapping_state = MappingRepository.load_mapping_state
build_vc_resolver = MappingRepository.build_vc_resolver



