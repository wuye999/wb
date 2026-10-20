# -*- coding: utf-8 -*-
"""商品价格表 D 列「实际售价」读取工具。

D 列由用户手工维护，可能是数值，也可能是公式（`=F2*2` / `=F2` / `=$F$2*2` /
`=ROUND(F2*1.3,2)` / 常量算式）。openpyxl 以 `data_only=False` 打开时只能拿到
公式字符串，因此这里做三层处理：

1. **受限求值**：只允许「单元格引用 + 四则运算 + 括号 + 数字常量」，
   剥掉 `ROUND(expr,n)` 外壳；绝不执行任意函数或裸 eval；
2. **缓存兜底**：求值失败时回退 Excel 保存的计算结果（`data_only=True`）；
3. **显式告警**：兜底与彻底失败都打 `[警告]`，避免价格静默变空。

口径：店铺整数价 = `math.floor(实际售价)`（见 `floor_price`）。
"""
import ast
import math
import operator
import re
import sys
from typing import Any, List, Optional, Tuple

from ..framework.xlsx_headers import (COL_ACTUAL_PRICE, COL_CN, find_col,
                                      read_headers, require_col)

# 单元格引用：F2 / $F$2 / F$2 / $F2（前后不能紧邻字母数字，避免误伤函数名/数字）
_CELL_RE = re.compile(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d{1,7})(?![0-9A-Z])")
# ROUND(表达式, 位数) —— 仅支持逗号分隔
_ROUND_RE = re.compile(r"^ROUND\((.+),\s*(\d+)\)$")
# 替换单元格引用后的安全表达式白名单
_SAFE_RE = re.compile(r"^[0-9eE+\-*/(). ]+$")
# 函数调用（SUM(...) / IF(...) 等）——出现即不支持，交给缓存兜底；普通括号不拦
_FUNC_RE = re.compile(r"[A-Z_][A-Z0-9_.]*\s*\(")

_OPS = {ast.Add: operator.add, ast.Sub: operator.sub,
        ast.Mult: operator.mul, ast.Div: operator.truediv,
        ast.USub: operator.neg, ast.UAdd: operator.pos}

# 商品价格表（BOSS_XLSX）列名 → 逻辑字段（按表头名定位，不依赖列位）
BOSS_COLUMNS = {
    "img": "产品图片",
    "sku": "卖家SKU",
    "cn": COL_CN,
    "dp": COL_ACTUAL_PRICE,
    "dim": "尺寸",
    "price": "最低售价",
    "prefix": "vendorCode前缀码",
    "stock_sku": "库存SKU",
}


def boss_columns(ws, path: str, required=("sku", "cn", "dp")):
    """解析商品价格表表头 → {逻辑字段: 列索引}；required 缺列抛 HeaderError。

    列位可任意调整（用户会插列/挪列），一律按表头名定位。
    """
    headers = read_headers(ws)
    cols = {}
    for key, name in BOSS_COLUMNS.items():
        cols[key] = (require_col(headers, name, "商品价格表", path)
                     if key in required else find_col(headers, name))
    return cols


def to_number(v: Any) -> Optional[float]:
    """数值 / 数字文本 → float；None、bool、中文文本等 → None。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def _safe_eval(expr: str) -> Optional[float]:
    """只对「数字 + - * / ( )」求值（AST 白名单，不执行任意代码）。"""
    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant) and not isinstance(node.value, bool) \
                and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
            return _OPS[type(node.op)](ev(node.left), ev(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in _OPS:
            return _OPS[type(node.op)](ev(node.operand))
        raise ValueError("unsupported node")

    try:
        return float(ev(ast.parse(expr, mode="eval")))
    except Exception:                      # noqa: BLE001 —— 语法/类型问题统一视为不可解析
        return None


def resolve_formula(formula: Any, ws, depth: int = 2) -> Optional[float]:
    """D 列公式 → 数值；不支持 / 无法求值 → None（调用方回退缓存值）。"""
    if not isinstance(formula, str):
        return to_number(formula)
    s = formula.strip()
    if not s.startswith("="):
        return to_number(s)
    s = s[1:].strip().upper().replace(" ", "").replace("_XLFN.", "")
    for _ in range(2):                     # 剥 ROUND(expr, n) 外壳
        m = _ROUND_RE.match(s)
        if not m:
            break
        s = m.group(1)
    # 只允许括号参与运算；出现「函数名(」式调用、跨表引用、区间 → 交给缓存兜底
    if _FUNC_RE.search(s) or "!" in s or ":" in s:
        return None

    import openpyxl                        # 延迟导入：仅公式路径需要
    from openpyxl.utils import column_index_from_string

    missing: List[str] = []

    def repl(m):
        col = column_index_from_string(m.group(1))
        val = ws.cell(row=int(m.group(2)), column=col).value
        num = to_number(val)
        if num is None and depth > 1 and isinstance(val, str) and val.strip().startswith("="):
            num = resolve_formula(val, ws, depth - 1)   # 引用单元格本身是公式
        if num is None:
            missing.append(m.group(0))
            return "0"
        return repr(num)

    expr = _CELL_RE.sub(repl, s)
    if missing or not _SAFE_RE.match(expr):
        return None
    return _safe_eval(expr)


def floor_price(v: Any) -> Optional[int]:
    """实际售价 → 店铺整数价（math.floor）；空 / 非法 / NaN / Inf → None。"""
    num = to_number(v)
    if num is None or math.isnan(num) or math.isinf(num):
        return None
    return math.floor(num)


class CachedDpReader:
    """懒加载 data_only 工作簿，按行号取 D 列缓存值（一次 load_boss 共享一份）。"""

    def __init__(self, path: str, sheet: str = "Sheet1", col: int = 4):
        self.path, self.sheet, self.col = path, sheet, col
        self._wb = None

    def get(self, row: int):
        if self._wb is None:
            import openpyxl
            self._wb = openpyxl.load_workbook(self.path, data_only=True)
        return self._wb[self.sheet].cell(row=row, column=self.col).value

    def close(self):
        if self._wb is not None:
            self._wb.close()
            self._wb = None


_WARNED = set()                            # 同一进程内同类问题只告警一次（shelve 会反复调用）


def warn_unresolved(path: str, by_cache: List[Tuple], missing: List[Tuple]) -> None:
    """公式不可解析时的统一告警（缓存兜底 / 彻底失败）。"""
    def _say(kind: str, rows: List[Tuple], msg: str):
        key = (kind, str(path), tuple(r[0] for r in rows))
        if key in _WARNED:
            return
        _WARNED.add(key)
        sample = "、".join(f"行{r}({sku})" for r, sku, _raw in rows[:5])
        tail = "…" if len(rows) > 5 else ""
        print(f"[警告] 商品价格表 D 列{msg}：{len(rows)} 行 → {sample}{tail}", file=sys.stderr)

    if by_cache:
        _say("cache", by_cache,
             "公式无法直接求值，已改用 Excel 缓存值（若刚改表且未保存，缓存可能过期）")
    if missing:
        _say("missing", missing,
             "公式无法解析且无可用缓存值，这些行实际售价为空（改价/上架会被跳过）")
