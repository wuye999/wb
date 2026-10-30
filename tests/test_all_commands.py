# -*- coding: utf-8 -*-
"""
wb_ops 全量命令与核心业务综合自动化测试套件

⚠ 日常不必全量跑（本套件真连平台，约 5 分钟）。按需测试请用选择器：
    python tests/run_tests.py --changed      # 只跑本次改动相关（推荐）
    python tests/run_tests.py --cmd appeals  # 只跑指定命令（新增/改命令时必跑）
    python tests/run_tests.py --help-smoke   # 最快回归：改动模块导入检查 + 全命令 --help
全量（跨层改动 / 发版）：
    python -m unittest tests/test_all_commands.py
"""
import os
import subprocess
import sys
import time
import unittest

PYTHON = sys.executable
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WB_SCRIPT = os.path.join(BASE_DIR, "wb.py")


class TestAllCommands(unittest.TestCase):

    def _run_cmd(self, args, expect_code=0):
        cmd = [PYTHON, WB_SCRIPT] + args
        env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
        res = subprocess.run(
            cmd,
            cwd=BASE_DIR,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=env,
        )
        if expect_code is not None:
            self.assertEqual(
                res.returncode,
                expect_code,
                f"Command failed: {' '.join(args)}\nStderr: {res.stderr}\nStdout: {res.stdout}",
            )
        return res


    # ---------------- 1. 全部 58 个命令的帮助与解析测试 ----------------
    def test_01_all_subcommand_helps(self):
        subcommands = [
            "shops", "fetch", "mapping", "mapping-import", "mapping-check",
            "mismatch-check", "review", "merge", "mapping-rename", "shops-mapping",
            "price", "price-wb", "price-bcs", "stock", "stock-wb", "stock-bcs", "trash",
            "trash-wb", "trash-bcs", "replicate", "import-shelve",
            "promo-apply", "promo-goods", "promo-gap", "discount", "discount-wb", "discount-scan", "discount-bcs",
            "dimension", "dims-check", "banned", "clean", "price-review",
            "orders", "questions", "questions-watch", "chat-watch", "chats", "appeals", "ratings", "ai-test",
            "mabang-orders",
            "mabang-forecast", "feishu-register", "feishu-register-dedup", "mabang-stock-register",
            "mabang-stock-daily", "feishu-vc-stats", "fbs-delivery", "order-stats", "mabang-process", "cookies-update",
            "cookie-refresh-wb",
            "daily", "schedule", "remote-wh", "shelve", "shelve-old"
        ]
        self.assertEqual(len(subcommands), 58)

        for subcmd in subcommands:
            with self.subTest(command=subcmd):
                res = self._run_cmd([subcmd, "--help"], expect_code=0)
                self.assertIn("usage: wb", res.stdout)

    def test_01b_vc_prefix_tail_format_mapping(self):
        """vendorCode 前缀正则须容忍「多变体序号」尾段（离线，纯正则断言）。

        背景：形如 `BCS-FZMG-334862442-2` / `BCS-ZZMX-191868961-15` 的卡（前缀后多一段 `-N`）
        此前不匹配 `config.VC_PREFIX_RE`（旧式要求「数字结尾 / 含斜杠」），导致
        `mapping.py` 前缀自动补录、`mapping_excel.py` 前缀兜底、`mapping_repo` 名称解析兜底
        **三处全部落空** → 每店恒定 55 条永久留在「未映射商品」（2026-09-21 实测：
        ZZMX 15 / FZMG 10 / GYDX 10 / MQXJ 10 / PPYD 10，三店完全同一批）。
        """
        import re
        from wb_ops import config

        rx = config.VC_PREFIX_RE
        # 兼容的合法格式（含新增的带序号尾段）
        valid = [
            ("BCS-CYQX-12345678", "CYQX"),                       # 标准
            ("BCS-CYQX-ozon-card-12345678", "CYQX"),             # 他人表 ozon-card 尾段
            ("BCS-QQNN-WRLINWI/1078999444", "QQNN"),             # 新供应商代码格式
            ("BCS-FZMG-334862442-2", "FZMG"),                    # ★ 多变体序号尾段（本次放宽）
            ("BCS-ZZMX-191868961-15", "ZZMX"),                   # ★
        ]
        for vc, expect in valid:
            with self.subTest(vc=vc):
                m = re.match(rx, vc)
                self.assertIsNotNone(m, f"应匹配成功: {vc}")
                self.assertEqual(m.group(1), expect)

        # 仍必须拒绝的非法格式（防止放宽后误收）
        invalid = [
            "BCS-qqnn-12345678",          # 小写前缀
            "BCS-1234-12345678",          # 数字前缀
            "BCS-QQN-12345678",           # 3 位前缀
            "BCS-QQNNN-12345678",         # 5 位前缀
            "BCS-QQNN-12345a",            # 数字尾段后跟字母
            "BCS-QQNN-12345-x",           # 序号尾段非数字
            "BCS-QQNN-",                  # 无尾段
            "ABC-QQNN-12345678",          # 非 BCS 开头
            "BCS-QQNN-abc-1",             # 首个尾段非数字
            "",
        ]
        for vc in invalid:
            with self.subTest(vc=vc):
                self.assertIsNone(re.match(rx, vc), f"应不匹配: {vc}")

    def test_01c_promo_goods_parse_offline(self):
        """promo-goods 离线解析断言（**不联网**）：查询串 / cmp 域请求头 / 供应商 UUID / 展平 / 本地反查。

        背景：cmp.wildberries.ru 广告推广接口 `GET /api/v1/adverts` 首次接入，
        被推广商品内嵌在 `content[].stocks.products[]`；抓包依据
        `api/网络请求/wb推广活动列表.har`（2026-09-23，status=[4,9,11]）。
        """
        from wb_ops.adapters import wb_ads_client as ads
        from wb_ops.services.discount import adverts

        # ① adverts 查询串：`[`/`]` 按 HAR 百分号编码、逗号保留
        url = ads.adverts_url(1, 100)
        self.assertIn("status=%5B4,9,11%5D", url)
        self.assertIn("bid_type=%5B1,2%5D", url)
        self.assertIn("type=%5B8,9%5D", url)
        self.assertIn("show_stocks=true", url)
        self.assertIn("page_number=1", url)
        self.assertIn("page_size=100", url)

        # ② cmp 域必需头：Authorization: Bearer <authorizev3>；x-supplierid 取自 cookie 的
        #    x-supplier-id-external（实测 credentials.json 三店均带，与 HAR 一致）
        shop = {"authorizev3": "JWT-TEST", "cookie": "a=1; x-supplier-id-external=UUID-XYZ; b=2"}
        h = ads.cmp_headers(shop)
        self.assertEqual(h["authorization"], "Bearer JWT-TEST")
        self.assertEqual(h["authorizev3"], "JWT-TEST")
        self.assertEqual(h["x-supplierid"], "UUID-XYZ")
        self.assertEqual(ads.supplier_uuid_from_cookie(shop), "UUID-XYZ")
        self.assertEqual(ads.supplier_uuid_from_cookie({"cookie": "a=1"}), "")  # 取不到留空 → 走兜底

        # ③ 展平：无商品的活动跳过；products_count>0 但明细为空 → 计入缺失清单（调用方打 [警告]）
        campaigns = [
            {"id": 40370920, "campaign_name": "c1", "status_id": 9, "payment_model": "cpc",
             "budget": 100, "create_date": "2026-09-23T06:01:11.811229+03:00", "products_count": 2,
             "stocks": {"products": [
                 {"nm": 1358933567, "name": "Переноска", "subject": {"name": "Переноски"},
                  "total_quantity_fbo": 0, "total_quantity_mp": 998},
                 {"nm": 1336777003, "name": "Пионы", "subject": {"id": 1156}, "total_quantity_fbo": 1}]}},
            {"id": 2, "campaign_name": "c2", "products_count": 3, "stocks": {"products": []}},
            {"id": 3, "campaign_name": "c3", "products_count": 0},
        ]
        flat = adverts.iter_campaign_products(campaigns)
        self.assertEqual(len(flat), 2)
        self.assertEqual([p["nm"] for _, p in flat], [1358933567, 1336777003])
        self.assertEqual(adverts.missing_detail_campaigns(campaigns), [2])

        # ④ --status 解析：显式给出走自定义，缺省回落后台默认视图 [4,9,11]
        class _Args:
            status = "9"
            no_cn = False
            page_size = 100
            limit = 0
            max_pages = 50

        self.assertEqual(adverts.build_options(_Args())["statuses"], (9,))
        _Args.status = ""
        opt = adverts.build_options(_Args())
        self.assertEqual(opt["statuses"], ads.STATUS_DEFAULT)
        self.assertTrue(opt["with_cn"])

        # ⑤ 本地真源反查：未收录 nmId → vc/cn 全空（调用方标注「本地真源未收录」，不联网核实）
        from wb_ops.storage.nm_resolver import NmResolver
        info = NmResolver(9352, with_cn=True).resolve(999999999999)
        self.assertEqual(info, {"vc": "", "cn": "", "src": ""})

    def test_01d_trash_wb_move_nms_offline(self):
        """trash 的 WB 原生移回收站（moveNmsToTrash）离线断言（**不联网**）。

        抓包依据 `api/网络请求/wb批量移至回收站.har`（2026-10-29）：
        body 为裸数组 `[{"nmID": int}]`，响应 `{error, errorText, additionalErrors}`，
        **不返回成功条数** ⇒ 只能按 nmID 逐条记账。
        """
        from wb_ops import common
        from wb_ops.services.replicate import trash_wb as T

        # ① 响应解析：仅列出的 nmID 为失败，未列出的视为成功
        self.assertEqual(T.parse_move_response({"error": False, "additionalErrors": {}}, [1, 2]), {})
        self.assertEqual(T.parse_move_response({"error": True, "errorText": "boom"}, [1, 2]),
                         {1: "boom", 2: "boom"})
        self.assertEqual(T.parse_move_response({"error": False, "additionalErrors": {"2": "有库存"}}, [1, 2]),
                         {2: "有库存"})
        self.assertEqual(
            T.parse_move_response({"error": False, "additionalErrors": {
                "errors": [{"params": {"1": "3"}, "trKey": "trDelivered"}]}}, [3, 4]),
            {3: "trDelivered"})
        # 无法归属到 nmID 的键 → 整批失败（保守，宁可报失败）
        self.assertEqual(T.parse_move_response({"error": False, "additionalErrors": {"globalErr": "x"}}, [5, 6]),
                         {5: "globalErr:x", 6: "globalErr:x"})

        # ② 分批与 body 形状（700 条 → 300/300/100）
        calls = []

        def fake_post(session, url, payload, allow_400_json=True):
            calls.append((url, payload))
            return {"error": False, "additionalErrors": {}}

        orig_ms, orig_post = T.wb_api.make_session, T.wb_api.request_post
        T.wb_api.make_session = lambda shop, rv=None: object()
        T.wb_api.request_post = fake_post
        try:
            r = T.move_nms({}, list(range(1, 701)), interval=0)
            self.assertEqual((r["ok"], r["fail"]), (700, 0))
            self.assertEqual([len(c[1]) for c in calls], [300, 300, 100])
            self.assertTrue(calls[0][0].endswith("/source/v2/moveNmsToTrash"))
            self.assertEqual(calls[0][1][0], {"nmID": 1})

            # ③ chunk 超限被夹到 300；重复 nmID 去重
            calls.clear()
            r = T.move_nms({}, [9, 9, 8, 8, 7], chunk=1000, interval=0)
            self.assertEqual([len(c[1]) for c in calls], [3])
            self.assertEqual([d["nmID"] for d in r["details"]], [9, 8, 7])

            # ④ 中途 cookie 失效：不外抛；已提交批次照实计 ok，剩余计失败
            calls.clear()

            def post_403(session, url, payload, allow_400_json=True):
                calls.append(payload)
                if len(calls) == 2:
                    raise common.CookieExpiredError("403：cookie 已失效")
                return {"error": False, "additionalErrors": {}}

            T.wb_api.request_post = post_403
            r = T.move_nms({}, list(range(1, 651)), interval=0)
            self.assertEqual((r["ok"], r["fail"]), (300, 350))
            self.assertTrue(r["aborted"])
            self.assertIn("cookie 失效", r["details"][-1]["msg"])

            # ⑤ 传输异常 → 该批全失败（不中断后续批）
            calls.clear()

            def post_boom(session, url, payload, allow_400_json=True):
                raise RuntimeError("conn reset")

            T.wb_api.request_post = post_boom
            r = T.move_nms({}, [11, 12], interval=0)
            self.assertEqual((r["ok"], r["fail"]), (0, 2))
            self.assertIn("conn reset", r["details"][0]["msg"])
        finally:
            T.wb_api.make_session, T.wb_api.request_post = orig_ms, orig_post

    def test_01e_default_warehouse_resolution(self):
        """仓库 ID 解析（WB 原生 offices/stores/short）：莫斯科 → 地区仓 → None（离线）。

        2026-10-29 起不再依赖 BCS，也不再读 known_warehouses.json。
        """
        from wb_ops.services.replicate import warehouse as wh

        orig = wh.list_stores
        try:
            # ① 命中「莫斯科仓库」→ 取 id（不是 warehouseId/officeId）
            wh._cache.clear()
            wh.list_stores = lambda sid: [
                {"id": 111, "name": "成都仓库", "warehouseName": "Chengdu"},
                {"id": 222, "name": "莫斯科仓库", "warehouseName": "Москва (Белая дача)"},
            ]
            self.assertEqual(wh.resolve_warehouse(9352), 222)

            # ② 没有莫斯科 → 回落俄罗斯地区仓（俄文 warehouseName 也能命中），且排除成都
            wh._cache.clear()
            wh.list_stores = lambda sid: [
                {"id": 333, "name": "成都仓库", "warehouseName": "Chengdu"},
                {"id": 444, "name": "二号仓", "warehouseName": "Подольск"},
            ]
            self.assertEqual(wh.resolve_warehouse(9352), 444)

            # ③ 全都不是主仓（只剩成都）→ None，且不缓存失败结果
            wh._cache.clear()
            wh.list_stores = lambda sid: [{"id": 333, "name": "成都仓库", "warehouseName": "Chengdu"}]
            self.assertIsNone(wh.resolve_warehouse(9352))
            self.assertEqual(wh._cache, {})

            # ④ 无仓库列表（无 WB 凭证 / 接口失败）→ None
            wh._cache.clear()
            wh.list_stores = lambda sid: []
            self.assertIsNone(wh.resolve_warehouse(9352))

            # ⑤ 成功结果命中进程内缓存
            wh._cache.clear()
            calls = []
            wh.list_stores = lambda sid: (calls.append(sid) or
                                          [{"id": 222, "name": "莫斯科仓库"}])
            wh.resolve_warehouse(9352)
            wh.resolve_warehouse(9352)
            self.assertEqual(len(calls), 1)
        finally:
            wh.list_stores = orig
            wh._cache.clear()

    def test_01g_warehouse_client_field_trap(self):
        """WB 仓库接口字段陷阱：必须取 `id`，不能取 `warehouseId`（后者是 officeId）（离线）。

        抓包依据 api/网络请求/wb获取仓库id.har（2026-10-29）：
        {"name":"莫斯科仓库","warehouseId":211,"id":1929635} —— 用错会清到错误仓库。
        """
        from wb_ops.adapters import wb_warehouse_client as wc

        capture = {}

        class _FakeSession:
            pass

        orig_ms, orig_req = wc.wb_api.make_session, wc.wb_api.request

        def fake_request(session, method, url, **kw):
            capture["url"] = url
            capture["params"] = kw.get("params")
            return {"error": False, "additionalErrors": {"errorCode": ""}, "data": [
                {"name": "成都仓库", "warehouseName": "382850-Chengdu", "warehouseId": 3006477,
                 "id": 1929632, "deliveryType": 1, "isDeleting": False},
                {"name": "莫斯科仓库", "warehouseName": "Москва (Белая дача)", "warehouseId": 211,
                 "id": 1929635, "deliveryType": 1, "isDeleting": False},
                {"name": "删除中仓库", "warehouseName": "x", "warehouseId": 9,
                 "id": 999999, "deliveryType": 1, "isDeleting": True},
            ]}

        try:
            wc.wb_api.make_session = lambda shop, rv=None: _FakeSession()
            wc.wb_api.request = fake_request
            rows = wc.get_stores({"shopId": 9352})

            self.assertTrue(capture["url"].endswith("/portal/offices/stores/short"))
            self.assertEqual(capture["params"], {"deliveryType": 1, "order": "asc"})
            # isDeleting 的条目被过滤
            self.assertEqual([r["id"] for r in rows], [1929632, 1929635])
            # 莫斯科仓取的是 id=1929635，而不是 warehouseId=211
            msk = next(r for r in rows if r["name"] == "莫斯科仓库")
            self.assertEqual(msk["id"], 1929635)
            self.assertEqual(msk["warehouseId"], 211)
            self.assertNotEqual(msk["id"], msk["warehouseId"])
        finally:
            wc.wb_api.make_session, wc.wb_api.request = orig_ms, orig_req

    def test_01h_mabang_mapping_header_based(self):
        """马帮本地映射按**表头名**定位：把列顺序彻底打乱后仍能取到正确 中文名/库存SKU（离线）。

        背景：旧实现硬编码 `r[1]/r[0]`（映射总表）与 `r[2]/r[7]`（商品价格表），
        用户插列/挪列后会静默读错——轻则全部判 NO_SKU，重则把订单替换成错误 SKU。
        """
        import os
        import shutil
        import tempfile

        from openpyxl import Workbook

        from wb_ops import config as cfg
        from wb_ops.services.order import mabang

        tmp = tempfile.mkdtemp(prefix="wb_hdr_")
        boss_xlsx = os.path.join(tmp, "boss.xlsx")
        map_xlsx = os.path.join(tmp, "map.xlsx")
        try:
            # 商品价格表：库存SKU 放第 2 列、产品中文名 放最后（旧实现 r[2]/r[7] 必错）
            wb = Workbook()
            ws = wb.active
            ws.append(["产品图片", "库存SKU", "卖家SKU", "实际售价", "尺寸",
                       "最低售价", "vendorCode前缀码", "产品中文名"])
            ws.append(["img", "TESTSKU-6161", "S-1", "=F2*2", "1*1*1/0.1",
                       "10", "AAAA", "测试商品"])
            wb.save(boss_xlsx)

            # 映射总表：vendorCode 在前、产品中文名 在后（旧实现 r[0]/r[1] 会读反）
            wb = Workbook()
            ws = wb.active
            ws.title = "映射总表"
            ws.append(["vendorCode", "产品中文名", "实际售价"])
            ws.append(["BCS-AAAA-123", "测试商品", 10])
            wb.save(map_xlsx)

            old_boss, old_map = cfg.BOSS_XLSX, cfg.MAPPING_XLSX
            cfg.BOSS_XLSX, cfg.MAPPING_XLSX = boss_xlsx, map_xlsx
            try:
                vc2cn, cn2sku = mabang.load_local_mapping()
            finally:
                cfg.BOSS_XLSX, cfg.MAPPING_XLSX = old_boss, old_map

            self.assertEqual(vc2cn.get("BCS-AAAA-123"), "测试商品")
            self.assertEqual(cn2sku.get("测试商品"), "TESTSKU-6161")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_01f_ops_channel_default(self):
        """默认通道映射：trash → wb（WB 原生）；price/stock → bcs（其 WB 通道走 price_wb/stock_wb）。"""
        from wb_ops.services.replicate.ops import _default_channel
        self.assertEqual(_default_channel("trash"), "wb")
        self.assertEqual(_default_channel("price"), "bcs")
        self.assertEqual(_default_channel("stock"), "bcs")

    # ---------------- 2. 账号与基础数据 ----------------
    def test_02_shops_list(self):
        res = self._run_cmd(["shops"], expect_code=0)
        self.assertIn("9352", res.stdout)
        self.assertIn("袁州", res.stdout)

    def test_03_fetch_cached(self):
        res = self._run_cmd(["fetch", "--shop-id", "9352", "--no-sync"], expect_code=0)
        self.assertIn("9352", res.stdout)
        self.assertIn("shop9352_products_all.json", res.stdout)

    # ---------------- 3. 折扣与改价业务 ----------------
    def test_04_discount_dry_run(self):
        res = self._run_cmd(["discount", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("BCS-HAAJ-248364237", res.stdout)

    def test_04b_discount_below(self):
        res = self._run_cmd(["discount", "--threshold", "55", "--below", "40", "--target", "50",
                             "--shops", "9352", "--limit", "5"], expect_code=0)
        self.assertIn("或", res.stdout)
        self.assertIn("<40", res.stdout)

    def test_04c_discount_upload_two_phase(self):
        """WB 改折扣两阶段提交（离线 mock，不触网/不写平台）。

        依据抓包 `api/网络请求/wb批量修改折扣+降价提示.har`：`upload/task` 必须
        先 `checkChange=true` 预检（只回弹窗标记、无 id），再 `checkChange=false`
        真正提交（回 `data.id`）。旧实现 URL 写死 checkChange=true，导致只做了预检、
        taskId 恒为 None，平台侧根本没落库。
        """
        from unittest import mock
        from wb_ops.adapters import wb_client as wc

        calls = []

        def fake_request(session, method, url, **kwargs):
            calls.append((method, url, kwargs.get("json")))
            if url.endswith("checkChange=true"):
                return {"data": {"priceModal": False, "quarantineModal": True}, "error": False, "errorText": ""}
            return {"data": {"id": 164469388, "alreadyExists": False}, "error": False, "errorText": ""}

        payload = [{"vendorCode": "BCS-QQNN-579331069", "nmID": 1579005260,
                    "discount": 49, "currencyIsoCode": "CNY"}]
        shop = {"shopId": 9352, "shopName": "袁州1", "cookie": "", "authorizev3": "t", "wb_seller_lk": "t"}
        client = wc.WBClient(shop, root_version="v1.113.3")

        with mock.patch.object(wc, "request", side_effect=fake_request):
            res = client.upload_batch_discount(payload)

        self.assertEqual([c[1].split("?")[-1] for c in calls],
                         ["checkChange=true", "checkChange=false"])
        self.assertEqual(calls[0][2], {"data": payload})
        self.assertEqual(calls[1][2], {"data": payload})
        self.assertTrue(res.success)
        self.assertEqual(res.task_id, 164469388)   # 有任务号 = 真落库
        self.assertTrue(res.quarantine_modal)
        self.assertFalse(res.price_modal)

        # precheck=False 时只提交一次，且仍必须带 checkChange=false
        calls.clear()
        with mock.patch.object(wc, "request", side_effect=fake_request):
            res2 = client.upload_batch_discount(payload, precheck=False)
        self.assertEqual(len(calls), 1)
        self.assertTrue(calls[0][1].endswith("checkChange=false"))
        self.assertEqual(res2.task_id, 164469388)
        self.assertFalse(res2.quarantine_modal)

    def test_05_discount_aliases(self):
        res1 = self._run_cmd(["discount-wb", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        res2 = self._run_cmd(["discount-scan", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("BCS-HAAJ-248364237", res1.stdout)
        self.assertIn("BCS-HAAJ-248364237", res2.stdout)

    def test_06_discount_bcs_dry_run(self):
        res = self._run_cmd(["discount-bcs", "--limit", "1", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_07_promo_apply_dry_run(self):
        res = self._run_cmd(["promo-apply", "--shops", "9352", "--days", "1"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_08_price_review_dry_run(self):
        res = self._run_cmd(["price-review", "--limit", "1", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    # ---------------- 4. Ops 一键操作 ----------------
    def test_09_price_dry_run(self):
        """price 默认通道 = WB 原生 dp-api 批量（2026-09-24 起切换，registry 路由 price_wb.run）。"""
        res = self._run_cmd(["price", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("WB 原生 dp-api", res.stdout)

    def test_09b_price_auto_quarantine_apply(self):
        """改价自动审核：隔离区按 nmID 精确匹配 + 跨域走 discount_svc 门面（离线 mock，不触网）。

        回归护栏：`ops_executor._auto_review_shop` 曾直接引用 discount 域私有模块
        `price_review` 却从未 import —— NameError 被 except Exception 吞掉，表现为
        「[自动审核] 店铺 X 失败: name 'price_review' is not defined」，自动审核**静默失效**
        （2026-09-21 实测）。现改为经 `discount_svc.apply_new_prices_by_nmids` 门面调用
        （REUSE_GUIDE 铁律 3：跨域只能经 `*_svc.py` 门面）。

        ⚠ 用例名刻意不含 review 字样：run_tests 选择器按「归一化子串 + 最长匹配」分派，
        `review` 与 `price` 等长时会把用例判给排在前面的 `review`，导致 `--cmd price` 漏测。
        """
        from unittest import mock
        from wb_ops.adapters import wb_client as wc
        from wb_ops.services.discount import price_review
        from wb_ops.services.discount_svc import discount_svc
        from wb_ops.services.replicate import ops_executor as oe

        quarantined = [{"id": 111, "nmID": 1001, "oldPrice": 198, "newPrice": 115},
                       {"id": 222, "nmID": 1002, "oldPrice": 158, "newPrice": 95},
                       {"id": 333, "nmID": 9999, "oldPrice": 100, "newPrice": 70}]  # 历史遗留待审
        posted = {}

        def fake_apply(session, ids):
            posted["ids"] = ids
            return {"error": False, "errorText": ""}

        shop = {"shopId": 9352, "shopName": "袁州1", "cookie": "",
                "authorizev3": "t", "wb_seller_lk": "t"}

        with mock.patch.object(wc, "make_session", return_value=object()), \
                mock.patch.object(price_review, "fetch_all_quarantine", return_value=quarantined), \
                mock.patch.object(price_review, "apply_prices", side_effect=fake_apply):

            # ① 门面：只审核 nm_ids 命中的待审项，不误审历史遗留（9999）
            r = discount_svc.apply_new_prices_by_nmids(shop, [1001, 1002])
            self.assertEqual(r["matched"], 2)
            self.assertEqual(r["applied"], 2)
            self.assertEqual(posted["ids"], [111, 222])

            # ② 调用方：改价后自动审核（原先此处恒 NameError → 返回 0，价格不生效）
            posted.clear()
            items = [{"nmID": 1001, "vc": "BCS-AAAA-1", "price": 115, "cur_price": 198.0},
                     {"nmID": 1002, "vc": "BCS-AAAA-2", "price": 95, "cur_price": 158.0},
                     {"nmID": 7777, "vc": "BCS-AAAA-3", "price": 500, "cur_price": 510.0}]  # 非 30-49.9%
            fake_cred = mock.Mock()
            fake_cred.wb_shop.return_value = shop
            with mock.patch.object(oe.credentials, "get", return_value=fake_cred):
                applied = oe._auto_review_shop(9352, items)
            self.assertEqual(applied, 2)          # 只审降价落 30-49.9% 的 2 个
            self.assertEqual(posted["ids"], [111, 222])

            # ③ 隔离区尚未生成（改价后平台异步入审查）→ 不报错、返回 0
            posted.clear()
            with mock.patch.object(price_review, "fetch_all_quarantine", return_value=[]):
                fake_cred2 = mock.Mock()
                fake_cred2.wb_shop.return_value = shop
                with mock.patch.object(oe.credentials, "get", return_value=fake_cred2):
                    self.assertEqual(oe._auto_review_shop(9352, items), 0)
            self.assertEqual(posted, {})

    def test_10_stock_dry_run(self):
        """stock 默认通道 = WB 原生在线接口（2026-09-24 起切换，registry 路由 stock_wb.run）。"""
        res = self._run_cmd(["stock", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("WB 原生在线接口", res.stdout)

    def test_09c_price_wb_dry_run(self):
        """price-wb dry-run：WB 原生 dp-api 通道（显式点名，不写平台）。
        用例名含归一化子串 pricewb → --cmd price-wb 精确命中；--cmd price 不会误捞。"""
        res = self._run_cmd(["price-wb", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("BCS-HAAJ-248364237", res.stdout)
        self.assertIn("WB 原生 dp-api", res.stdout)

    def test_09d_price_bcs_dry_run(self):
        """price-bcs 备选通道：BCS price/batch（dry-run 预览，无「WB 原生」字样）。"""
        res = self._run_cmd(["price-bcs", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertNotIn("WB 原生 dp-api", res.stdout)

    def test_09e_price_wb_offline(self):
        """price_wb 离线断言（**不触网、不写平台**）。

        依据 HAR `api/网络请求/wb在线批量修改价格.har`：
        - 提交条目 {"vendorCode","nmID","price","currencyIsoCode":"CNY"}（只改价无 discount 字段）
        - checkChange=true 预检（可能返回 priceModal/quarantineModal）→ checkChange=false 落库
          =「自动确认提交」（用户确认的语义，与改折扣两阶段同机制）
        """
        from unittest import mock
        from argparse import Namespace
        from wb_ops.adapters import wb_client as wc
        from wb_ops.services.replicate import price_wb

        shop = {"shopId": 9352, "shopName": "袁州1", "cookie": "",
                "authorizev3": "t", "wb_seller_lk": "t"}

        # ① 条目映射：只改价（无 discount 字段）/ --discount 显式给出（含 0）才带
        it = {"vc": "BCS-X-1", "nmID": 123, "price": 150, "cur_price": 198.0,
              "orig_zero": False, "discount": None, "clubDiscount": None, "cn": "测试"}
        p0 = price_wb._build_payload(it, None)
        self.assertEqual(p0, {"vendorCode": "BCS-X-1", "nmID": 123,
                              "price": 150, "currencyIsoCode": "CNY"})
        self.assertEqual(price_wb._build_payload(it, 30)["discount"], 30)
        self.assertEqual(price_wb._build_payload(it, 0)["discount"], 0)   # 显式 0 也照传

        # ② exec_shop：checkChange 两阶段 + 预检 modal=true 仍自动确认落库 + 分批
        calls = []

        def fake_request(session, method, url, **kwargs):
            calls.append((url.split("?")[-1], kwargs.get("json")))
            if url.endswith("checkChange=true"):
                return {"data": {"priceModal": True, "quarantineModal": True}, "error": False, "errorText": ""}
            return {"data": {"id": 166628902, "alreadyExists": False}, "error": False, "errorText": ""}

        items = [{"vc": f"BCS-WB-{i}", "nmID": 1000 + i, "price": 150, "cur_price": 198.0,
                  "orig_zero": False, "discount": None, "clubDiscount": None, "cn": "测试"}
                 for i in range(4)]
        plan = {"shopId": 9352, "items": items}
        pargs = Namespace(chunk=2, interval=0, discount=None)
        client = wc.WBClient(shop, root_version="v1.113.3")
        results, ts = [], "2026-09-24 12:00:00"
        with mock.patch.object(wc, "request", side_effect=fake_request):
            ok, fail = price_wb.exec_shop(client, plan, pargs, results, ts)
        self.assertEqual((ok, fail), (4, 0))
        self.assertEqual([c[0] for c in calls],
                         ["checkChange=true", "checkChange=false", "checkChange=true", "checkChange=false"])
        self.assertEqual(calls[0][1], {"data": [
            {"vendorCode": "BCS-WB-0", "nmID": 1000, "price": 150, "currencyIsoCode": "CNY"},
            {"vendorCode": "BCS-WB-1", "nmID": 1001, "price": 150, "currencyIsoCode": "CNY"}]})
        self.assertEqual(len(results), 4)
        self.assertTrue(all(r[2] == "price_wb" for r in results))

        # ③ apply_prices 门面：chunk 分批 + alreadyExists 计成功 + task_ids 收集
        def fake_request_exists(session, method, url, **kwargs):
            if url.endswith("checkChange=true"):
                return {"data": {"priceModal": False, "quarantineModal": False}, "error": False, "errorText": ""}
            return {"data": {"id": 166628902, "alreadyExists": True}, "error": False, "errorText": ""}

        witems = [{"vendorCode": f"BCS-WB-{i}", "nmID": 2000 + i, "price": 88} for i in range(3)]
        calls.clear()
        with mock.patch.object(wc, "request", side_effect=fake_request_exists):
            r = price_wb.apply_prices(shop, witems, root_version="v1.113.3", chunk=2, interval=0)
        self.assertEqual((r["ok"], r["fail"]), (3, 0))
        self.assertEqual(r["task_ids"], [166628902, 166628902])
        self.assertEqual(len(r["details"]), 3)
        self.assertTrue(all(d["ok"] for d in r["details"]))

        # ④ ≤原价/2 预拦截：price=99 ≤ cur/2=99 → 剔除不提交
        bad_items = [{"vc": "BCS-WB-9", "nmID": 9999, "price": 99, "cur_price": 198.0,
                      "orig_zero": False, "discount": None, "clubDiscount": None, "cn": "测试"}]
        plan_bad = {"shopId": 9352, "items": bad_items}
        calls.clear()
        with mock.patch.object(wc, "request", side_effect=fake_request):
            ok2, fail2 = price_wb.exec_shop(client, plan_bad, pargs, [], ts)
        self.assertEqual((ok2, fail2), (0, 0))
        self.assertEqual(calls, [])          # 全部被拦截 → 不发任何请求

    def test_09f_price_boss_dp_formula_unit(self):
        """price 目标价：商品价格表 D 列「实际售价」公式解析 + floor 口径（离线）。

        覆盖 2026-10-20 口径变更：D 列由「双倍售价」改为「实际售价」，
        不再假设 =F*2；需支持 =F2*2 / =F2 / =$F$2*2 / =ROUND(...) / 纯数值，
        无法解析时回退 Excel 缓存值，两者皆无则 None 并告警。
        """
        import contextlib
        import io as _io
        import os
        import shutil
        import tempfile
        from unittest import mock

        import openpyxl
        from wb_ops import config
        from wb_ops.services.replicate import ops_plan
        from wb_ops.storage.mapping_repo import MappingRepository

        header = ["产品图片", "卖家SKU", "产品中文名", "实际售价", "尺寸",
                  "最低售价", "vendorCode前缀码", "库存SKU"]
        rows = [
            ["", "SKU-1", "旧格式", "=F2*2", "10*10*10/0.1", 40.25, "AAAA", "s1"],
            ["", "SKU-2", "无乘数", "=F3", "10*10*10/0.1", 41.525, "BBBB", "s2"],
            ["", "SKU-3", "绝对引用", "=$F$4*2", "10*10*10/0.1", 39.775, "CCCC", "s3"],
            ["", "SKU-4", "ROUND", "=ROUND(F5*1.3,2)", "10*10*10/0.1", 46.2, "DDDD", "s4"],
            ["", "SKU-5", "走缓存", "=VLOOKUP(A2,A:F,6,0)", "10*10*10/0.1", 50, "EEEE", "s5"],
            ["", "SKU-6", "纯数值", 100.75, "10*10*10/0.1", 60, "FFFF", "s6"],
            ["", "SKU-7", "无缓存", "=IF(A2>0,1,2)", "10*10*10/0.1", 70, "GGGG", "s7"],
        ]
        wb_formula = openpyxl.Workbook()
        ws = wb_formula.active
        ws.title = "Sheet1"
        ws.append(header)
        for r in rows:
            ws.append(r)

        wb_cached = openpyxl.Workbook()
        ws2 = wb_cached.active
        ws2.title = "Sheet1"
        ws2.append(header)
        for r in rows:
            ws2.append([r[0], r[1], r[2], None, r[4], r[5], r[6], r[7]])
        ws2.cell(2, 4).value = 80.5      # =F2*2
        ws2.cell(3, 4).value = 41.525    # =F3
        ws2.cell(4, 4).value = 79.55     # =$F$4*2
        ws2.cell(5, 4).value = 60.06     # =ROUND(...)
        ws2.cell(6, 4).value = 88.88     # 仅走缓存兜底（故意与 F5*2 不同，便于区分来源）
        # 第 8 行（=IF...）不写缓存 → 应得 None 并告警

        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "boss.xlsx")
            wb_formula.save(path)
            err = _io.StringIO()
            with mock.patch("openpyxl.load_workbook",
                            side_effect=lambda p, data_only=False, **kw:
                            wb_cached if data_only else wb_formula), \
                    mock.patch.object(config, "BOSS_XLSX", path), \
                    contextlib.redirect_stderr(err):
                boss = MappingRepository.load_boss()

            self.assertEqual([b["dp"] for b in boss],
                             [80.5, 41.525, 79.55, 60.06, 88.88, 100.75, None])
            self.assertEqual([b["floor"] for b in boss],
                             [80, 41, 79, 60, 88, 100, None])
            self.assertIn("缓存", err.getvalue())          # 回退缓存 → 可见
            self.assertIn("无法解析", err.getvalue())      # 彻底失败 → 可见

            # 映射表快照 dp 兜底也必须是 floor（修 int("80.5") 崩溃）
            self.assertEqual(
                ops_plan.target_price("VC-X", {"VC-X": {"cn": "无", "dp": 88.9}}, [], None), 88)
            self.assertIsNone(
                ops_plan.target_price("VC-Y", {"VC-Y": {"cn": "无", "dp": "bad"}}, [], None))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_10b_stock_bcs_dry_run(self):
        """stock-bcs 备选通道：BCS stock/batchSetByChrtIdsBatch（dry-run 预览）。"""
        res = self._run_cmd(["stock-bcs", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_10c_stock_wb_dry_run(self):
        """stock-wb dry-run：WB 原生在线接口通道（快照解析，不写平台）。
        用例名含归一化子串 stockwb → --cmd stock-wb 精确命中；
        --cmd stock 时最长匹配不会误捞本用例（归一化 stockwb ⊃ stock 反向不成立）。"""
        res = self._run_cmd(["stock-wb", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("BCS-HAAJ-248364237", res.stdout)
        self.assertIn("WB 原生在线接口", res.stdout)

    def test_10d_stock_wb_offline(self):
        """stock-wb 离线断言（**不触网、不写平台**）：

        依据 HAR `api/网络请求/wb在线加载修改库存.har`：
        - 修改端点 POST .../api/v3/portal/stocks/{warehouseId}，body {"data":[{chrtId,amount}]}
        - 数组批量（探针 _scratch/probe_stock_wb_capacity.py 实测 1000 条 error=false）
        - 查询端点 GET 同域 portal/stocks?order=asc[&stores=][&search=]（data.next 游标分页）
        """
        from unittest import mock
        from wb_ops.adapters import wb_stock_client as wsc
        from wb_ops.services.replicate import stock_wb

        shop = {"shopId": 9352, "shopName": "袁州1", "cookie": "",
                "authorizev3": "t", "wb_seller_lk": "t"}

        # ① post_stocks：URL 以 warehouseId 结尾、body 结构与 allow_400_json 默认开启
        posts = []

        def fake_request_post(session, url, payload, allow_400_json=False):
            posts.append((url, payload, allow_400_json))
            return {"error": False, "errorText": "", "data": {}}

        with mock.patch.object(wsc.wb_api, "request_post", side_effect=fake_request_post):
            r = wsc.post_stocks(object(), 1929635, [{"chrtId": 2102520529, "amount": 30}])
        self.assertTrue(posts[0][0].endswith("/api/v3/portal/stocks/1929635"))
        self.assertEqual(posts[0][1], {"data": [{"chrtId": 2102520529, "amount": 30}]})
        self.assertTrue(posts[0][2])          # WB 400 也带 JSON 错误体 → allow_400_json 默认 True
        self.assertFalse(r["error"])

        # ② set_stock 门面：201 条按 chunk=200 分 200+1 两批；明细 201 条全成功
        items = [{"chrtId": 1000 + i, "amount": i} for i in range(201)]
        posts.clear()
        with mock.patch.object(wsc.wb_api, "request_post", side_effect=fake_request_post), \
                mock.patch.object(wsc.wb_api, "make_session", return_value=object()):
            r2 = stock_wb.set_stock(shop, 1929635, items, chunk=200, interval=0)
        self.assertEqual([len(p[1]["data"]) for p in posts], [200, 1])
        self.assertEqual((r2["ok"], r2["fail"]), (201, 0))
        self.assertEqual(len(r2["details"]), 201)
        self.assertTrue(all(d["ok"] for d in r2["details"]))

        # ③ chunk 上限夹取：chunk=5000 传入门面 → 实际按 1000 分批（1001 条 → 1000+1）
        items_1001 = [{"chrtId": 2000 + i, "amount": i} for i in range(1001)]
        posts.clear()
        with mock.patch.object(wsc.wb_api, "request_post", side_effect=fake_request_post), \
                mock.patch.object(wsc.wb_api, "make_session", return_value=object()):
            r3 = stock_wb.set_stock(shop, 1929635, items_1001, chunk=5000, interval=0)
        self.assertEqual([len(p[1]["data"]) for p in posts], [1000, 1])
        self.assertEqual(r3["ok"], 1001)

        # ④ 失败路径：error=true → 全部计失败，msg 取 errorText
        def fake_post_fail(session, url, payload, allow_400_json=False):
            return {"error": True, "errorText": "Недостаточно прав", "data": {}}

        with mock.patch.object(wsc.wb_api, "request_post", side_effect=fake_post_fail), \
                mock.patch.object(wsc.wb_api, "make_session", return_value=object()):
            r4 = stock_wb.set_stock(shop, 1929635, items[:3], chunk=200, interval=0)
        self.assertEqual((r4["ok"], r4["fail"]), (0, 3))
        self.assertEqual(r4["details"][0]["msg"], "Недостаточно прав")

        # ⑤ get_stocks：stores/search 参数透传 + data.next 游标翻页 + 空游标终止
        pages = [
            {"data": {"next": "CUR1",
                      "stocks": [{"article": "BCS-A-1", "chrtId": 11, "amount": 1}]}},
            {"data": {"next": "",
                      "stocks": [{"article": "BCS-B-2", "chrtId": 22, "amount": 0}]}},
        ]
        gets = []

        def fake_request(session, method, url, **kwargs):
            gets.append(kwargs.get("params") or {})
            return pages[len(gets) - 1]

        with mock.patch.object(wsc.wb_api, "request", side_effect=fake_request):
            got = wsc.get_stocks(object(), store_id=1929635, search="BCS-A-1")
        self.assertEqual([g["chrtId"] for g in got], [11, 22])
        self.assertEqual(gets[0].get("stores"), 1929635)
        self.assertEqual(gets[0].get("search"), "BCS-A-1")
        self.assertNotIn("next", gets[0])     # 首页不带 next
        self.assertEqual(gets[1].get("next"), "CUR1")
        self.assertEqual(gets[1].get("search"), "BCS-A-1")  # search 每页透传

    def test_11_trash_dry_run(self):
        res = self._run_cmd(["trash", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("WB 原生 moveNmsToTrash", res.stdout)   # 默认通道 = WB 原生

    def test_11b_trash_bcs_dry_run(self):
        res = self._run_cmd(["trash-bcs", "--vc", "BCS-HAAJ-248364237", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("BCS clean/removeToTrash", res.stdout)  # 备用通道 = BCS

    # ---------------- 5. 商品搬家与清理 ----------------
    def test_12_replicate_dry_run(self):
        res = self._run_cmd(["replicate", "--vc", "BCS-HAAJ-248364237", "--limit", "1"], expect_code=0)
        self.assertIn("vendorCode", res.stdout)

    def test_13_dimension_dry_run(self):
        res = self._run_cmd(["dimension", "--vc", "BCS-HAAJ-248364237", "--shops", "9352", "--dims", "8*14*26/0.3"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_14_dims_check(self):
        res = self._run_cmd(["dims-check", "--limit", "1", "--shops", "9352"], expect_code=0)
        self.assertIn("袁州1", res.stdout)

    def test_15_banned_dry_run(self):
        res = self._run_cmd(["banned", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_16_clean_dry_run(self):
        res = self._run_cmd(["clean", "--target", "basket", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_17_remote_wh_dry_run(self):
        res = self._run_cmd(["remote-wh", "--shops", "9352"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    # ---------------- 6. 映射与工作台 ----------------
    def test_18_shops_mapping(self):
        res = self._run_cmd(["shops-mapping", "--shop-id", "9352"], expect_code=0)
        self.assertIn("shop_9352", res.stdout)

    def test_19_mapping_check(self):
        res = self._run_cmd(["mapping-check", "--tol", "5"], expect_code=0)
        self.assertIn("映射表核查工作台.html", res.stdout)

    def test_20_mismatch_check(self):
        res = self._run_cmd(["mismatch-check", "--days", "1"], expect_code=0)
        self.assertIn("货不对板筛查工作台.html", res.stdout)

    def test_21_review(self):
        res = self._run_cmd(["review"], expect_code=0)
        self.assertIn("待审核工作台已生成", res.stdout)

    # ---------------- 7. 订单与马帮 ----------------
    def test_22_orders(self):
        res = self._run_cmd(["orders", "--no-sync", "--days", "1", "--shops", "9352"], expect_code=0)
        self.assertIn("订单", res.stdout)

    def test_23_mabang_orders_dry_run(self):
        res = self._run_cmd(["mabang-orders", "--days", "1"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_23b_mabang_orders_csv_fallback_unit(self):
        """离线单测：中文报告名被环境按文件名拦截时，自动回退 ASCII 别名且不影响主流程。

        同时校验：两种命名都失败时 write_csv 返回 ""（不抛异常），
        避免「报告写不出来 → 整个订单链路中止」。
        """
        import csv as _csv
        import shutil
        import tempfile
        from unittest import mock
        from wb_ops.services.order import mabang

        rec = {"status": "REPLACE", "order_id": "421000001", "platform_order_id": "P-1",
               "shop_bcs": "袁州1", "shop": "9352", "paid_time": "2026-10-17 09:00:00",
               "vc": "BCS-AAAA-100000001", "cn_name": "测试品", "matched_sku": "SKU-1",
               "expected_sku": "SKU-1", "note": "字符串一致"}
        real_open = open
        tmp = tempfile.mkdtemp()
        try:
            # ① 仅拦中文名 → 应回退到 ASCII 别名并正常落盘
            def block_cn(path, *a, **kw):
                if "马帮订单匹配" in str(path):
                    raise PermissionError(13, "Permission denied", str(path))
                return real_open(path, *a, **kw)

            with mock.patch("builtins.open", side_effect=block_cn), \
                    mock.patch.object(mabang.config, "LOG_DIR", tmp):
                path = mabang.write_csv([rec])
            self.assertTrue(path, "中文名被拦后应回退成功并返回实际路径")
            self.assertIn("mabang_order_match_", os.path.basename(path))
            with real_open(path, encoding="utf-8-sig") as f:
                rows = list(_csv.DictReader(f))
            self.assertEqual(rows[0]["订单ID"], "421000001")
            self.assertEqual(rows[0]["状态"], "REPLACE")

            # ② 两种命名都被拦 → 返回 ""（不抛异常，主链路继续）
            def block_all(path, *a, **kw):
                raise PermissionError(13, "Permission denied", str(path))

            with mock.patch("builtins.open", side_effect=block_all), \
                    mock.patch.object(mabang.config, "LOG_DIR", tmp):
                self.assertEqual(mabang.write_csv([rec]), "",
                                 "全部失败时应静默降级返回空路径")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_24_mabang_forecast_check(self):
        res = self._run_cmd(["mabang-forecast", "--check"], expect_code=0)
        self.assertIn("预报批次", res.stdout)

    def test_24b_mabang_forecast_check_overview_unit(self):
        """离线单测：--check 四 tab 概览（mock get_forecast_list，按 status 分派）。

        tab → status：1 待预报 / 3 预报成功 / 99 预报失败 / 5 历史预报。
        """
        import contextlib
        import io as _io
        from types import SimpleNamespace
        from unittest import mock
        from wb_ops.services.order import mabang

        def fake_list(cred, status=1, rows_per_page=200, page=1, **kw):
            rows = {"1": [], "3": [{"batchNo": "B3", "shopId": "主号7", "status": 3, "total": 1,
                                    "successNum": 1, "failNum": 0, "handoverNumber": "WB-GI-1",
                                    "createTime": "2026-09-30 09:52:14",
                                    "employeeName": "1503", "forecastLogisticsName": "WB线上",
                                    "orderWeight": "1.0000"}],
                    "99": [], "5": [{"batchNo": "B5", "shopId": "子龙", "status": 3, "total": 2,
                                     "successNum": 2, "failNum": 0}]}[str(status)]
            stats = {"waitTotal": 0, "succesTotal": 1, "failTotal": 0, "historyTotal": 1,
                     "pageTotal": len(rows), "pageFrom": 1 if rows else 0,
                     "pageTo": len(rows), "pageCur": 1 if rows else 0,
                     "pageCount": 1 if rows else 0}
            return rows, stats

        args = SimpleNamespace(check=True, status=None, pages=1)
        buf = _io.StringIO()
        with mock.patch.object(mabang, "get_forecast_list", side_effect=fake_list), \
                mock.patch.object(mabang, "_mabang_cred", return_value={}), \
                contextlib.redirect_stdout(buf):
            code = mabang.run_forecast(args)
        out = buf.getvalue()
        self.assertEqual(code, 0)
        for label in ("待预报", "预报成功", "预报失败", "历史预报"):
            self.assertIn(label, out, f"四 tab 概览应包含「{label}」段")
        self.assertIn("B3", out)

    def test_25_mabang_process_dry_run(self):
        res = self._run_cmd(["mabang-process", "--days", "1"], expect_code=0)
        self.assertIn("步骤", res.stdout)

    @staticmethod
    def _fb(found=None, stats=None, missing=None, pages=1):
        """构造 find_batches 返回值（离线单测用）。"""
        found = found or {}
        return {"found": found, "missing": set(missing or []),
                "all_found": not (missing or []), "pages_fetched": pages,
                "stats": stats or {}}

    def test_25b_mabang_process_upload_wait_unit(self):
        """离线单测：正向判据 —— 命中 status=3「预报成功」即完成 + 店铺映射提取。

        预扫描 status=1 命中（上传中 status=2，建立映射）；
        第 1 轮 status=3 未命中、第 2 轮 status=3 命中 ⇒ done=True。
        并校验新节奏：status=1 仅预扫描一次（不再每轮查）。
        """
        from unittest import mock
        from wb_ops.services.order import mabang

        row = {"batchNo": "B1", "shopId": "子龙主2", "status": 3,
               "total": 3, "successNum": 3, "failNum": 0, "handoverNumber": "WB-GI-1"}
        up1 = {"batchNo": "B1", "shopId": "子龙主2", "status": 2,
               "total": 3, "successNum": 0, "failNum": 0}
        rounds, calls = {"n": 0}, []

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            calls.append(status)
            if status == 3:
                rounds["n"] += 1
                if rounds["n"] >= 2:
                    return self._fb(found={"B1": row}, stats={"succesTotal": 1})
                return self._fb(stats={"succesTotal": 0, "waitTotal": 1})
            if status == 1:
                return self._fb(found={"B1": up1}, stats={"waitTotal": 1})
            return self._fb()

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None):
            res = mabang.wait_upload_done(None, ["B1"], timeout=60, interval=1)

        self.assertTrue(res["done"], "命中 status=3 后应判定上传成功完成")
        self.assertEqual(res["success"], {"B1"})
        self.assertEqual(res["store_of"].get("B1"), "子龙主2", "应记录 batchNo→shopId 映射")
        self.assertTrue(res["rows"].get("B1"), "应保留批次快照")
        self.assertEqual(res["failed"], set())
        self.assertEqual(calls.count(1), 1, "status=1 应只在预扫描查一次（每轮仅查 status=3）")
        self.assertGreaterEqual(calls.count(3), 1, "主循环应轮询 status=3")

    def test_25c_mabang_process_upload_wait_timeout_unit(self):
        """离线单测：轮询超时且批次仍在待预报 ⇒ done=False，返回未完成集合与映射。"""
        from unittest import mock
        from wb_ops.services.order import mabang

        row = {"batchNo": "B9", "shopId": "子龙主2（2）", "status": 1,
               "total": 2, "successNum": 0, "failNum": 0}
        clock = {"v": 0}

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            return self._fb(found={"B9": row}, stats={"waitTotal": 1}) if status == 1 else self._fb()

        def fake_time():
            clock["v"] += 10      # 每次读表前进 10s ⇒ 第一轮即越过 timeout
            return clock["v"]

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None), \
                mock.patch.object(mabang.time, "time", fake_time):
            res = mabang.wait_upload_done(None, ["B9"], timeout=5, interval=1)

        self.assertFalse(res["done"], "批次始终未成功时应判定超时")
        self.assertEqual(res["still"], {"B9"})
        self.assertEqual(res["failed"], set())
        self.assertEqual(res["store_of"].get("B9"), "子龙主2（2）")

    def test_25d_mabang_process_forecast_failed_unit(self):
        """离线单测：超时后经 status=99 定性为「预报失败」⇒ failed 集合（不交运）。"""
        from unittest import mock
        from wb_ops.services.order import mabang

        fr = {"batchNo": "B7", "shopId": "子龙主2", "status": 99,
              "total": 2, "successNum": 0, "failNum": 2}
        clock = {"v": 0}

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            return self._fb(found={"B7": fr}, stats={"failTotal": 1}) if status == 99 else self._fb()

        def fake_time():
            clock["v"] += 10
            return clock["v"]

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None), \
                mock.patch.object(mabang.time, "time", fake_time):
            res = mabang.wait_upload_done(None, ["B7"], timeout=5, interval=1)

        self.assertFalse(res["done"])
        self.assertEqual(res["failed"], {"B7"})
        self.assertEqual(res["still"], set(), "定性命中失败的批次不应再计入未完成")
        self.assertEqual(res["store_of"].get("B7"), "子龙主2")

    def test_25e_mabang_process_history_success_unit(self):
        """离线单测：批次已归档到 status=5「历史预报」亦算成功（防永不命中而空等超时）。

        同时校验 status=5 兜底的**限频**：连续 3 轮 status=3 未命中才触发一次。
        """
        from unittest import mock
        from wb_ops.services.order import mabang

        row = {"batchNo": "B1", "shopId": "袁州3", "status": 3, "isHistoryStatus": "5",
               "total": 1, "successNum": 1, "failNum": 0}
        calls = []

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            calls.append(status)
            return self._fb(found={"B1": row}, stats={"historyTotal": 1}) if status == 5 else self._fb()

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None):
            res = mabang.wait_upload_done(None, ["B1"], timeout=60, interval=1)

        self.assertTrue(res["done"])
        self.assertEqual(res["success"], {"B1"})
        self.assertEqual(res["store_of"].get("B1"), "袁州3")
        self.assertEqual(calls.count(5), 1, "status=5 兜底应限频（3 轮未命中触发一次）")
        self.assertEqual(calls.count(1), 1, "status=1 仅预扫描一次")

    def test_25g_mabang_process_prescan_map_timeout_unit(self):
        """离线单测：预扫描建立店铺映射；status=3 恒不命中 ⇒ 超时时映射仍齐全（避免整体跳过）。"""
        from unittest import mock
        from wb_ops.services.order import mabang

        row = {"batchNo": "B1", "shopId": "子龙主2", "status": 2,
               "total": 3, "successNum": 0, "failNum": 0}
        calls, clock = [], {"v": 0}

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            calls.append(status)
            return self._fb(found={"B1": row}, stats={"waitTotal": 1}) if status == 1 else self._fb()

        def fake_time():
            clock["v"] += 10
            return clock["v"]

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None), \
                mock.patch.object(mabang.time, "time", fake_time):
            res = mabang.wait_upload_done(None, ["B1"], timeout=5, interval=1)

        self.assertFalse(res["done"])
        self.assertEqual(res["still"], {"B1"})
        self.assertEqual(res["store_of"].get("B1"), "子龙主2", "预扫描应建立店铺映射")
        self.assertEqual(calls.count(1), 1, "映射已齐全 ⇒ 超时补查 status=1 应跳过")

    def test_25h_mabang_process_fallback5_throttled_unit(self):
        """离线单测：status=3 连续未命中时 status=5 兜底触发，且全程只触发一次（限频）。"""
        from unittest import mock
        from wb_ops.services.order import mabang

        calls = []
        fake_time = type("T", (), {"t": 0})()

        def _time():
            fake_time.t += 1      # 每轮推进 1s：miss≥3 与「过半超时」同期触发 ⇒ 只查一次
            return fake_time.t

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            calls.append(status)
            return self._fb(stats={"waitTotal": 1})     # 全部不命中

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None), \
                mock.patch.object(mabang.time, "time", _time):
            res = mabang.wait_upload_done(None, ["B1"], timeout=8, interval=1)

        self.assertFalse(res["done"])
        self.assertEqual(calls.count(5), 1, "status=5 兜底应限频为一次")

    def test_25j_mabang_process_fail_early_unit(self):
        """离线单测：failTotal 上涨 ⇒ 立刻核实 status=99 定性失败并早退，不空等到超时。

        复现 2026-10-17 袁州1 29 单场景：批次落在 status=99，但轮询只看 status=3，
        旧逻辑要等满 900s 才定性（白等 40+ 轮）；修复后应在 failTotal 一涨就收口。
        """
        from unittest import mock
        from wb_ops.services.order import mabang

        fr = {"batchNo": "B1", "shopId": "袁州1", "status": 99,
              "total": 29, "successNum": 0, "failNum": 29}
        state = {"r3": 0, "fail": 0}
        calls = []

        def fake_find(cred, batch_nos, status, rows_per_page=200, max_pages=50):
            calls.append(status)
            if status == 3:
                state["r3"] += 1
                if state["r3"] >= 2:          # 第 2 轮起平台把该批次计入失败
                    state["fail"] = 1
                return self._fb(stats={"failTotal": state["fail"], "waitTotal": 0})
            if status == 99:
                return self._fb(found={"B1": fr}, stats={"failTotal": state["fail"]})
            return self._fb(stats={"failTotal": 0})      # status=1 预扫描

        clock = {"v": 0}

        def fake_time():
            clock["v"] += 20
            return clock["v"]

        with mock.patch.object(mabang, "find_batches", side_effect=fake_find), \
                mock.patch.object(mabang.time, "sleep", lambda *_: None), \
                mock.patch.object(mabang.time, "time", fake_time):
            res = mabang.wait_upload_done(None, ["B1"], timeout=900, interval=20)

        self.assertFalse(res["done"])
        self.assertEqual(res["failed"], {"B1"})
        self.assertEqual(res["still"], set(), "已定性失败的批次不算未完成")
        self.assertEqual(res["reason"], "failed")
        self.assertEqual(calls.count(99), 1, "应对 status=99 核实一次")
        self.assertLessEqual(calls.count(3), 3,
                             "failTotal 上涨后应早退，不应跑满 900s/45 轮")

    def test_25k_mabang_process_replace_retry_unit(self):
        """离线单测：更换失败 → REPLACE_FAIL（不进预报/交运）+ 网络类异常重试1次 + 失败专表。

        覆盖 2026-10-19 事故：9 单 `Read timed out` 更换失败后仍被预报/交运。
        """
        import csv as _csv
        import shutil
        import tempfile

        import requests
        from unittest import mock
        from wb_ops.services.order import mabang

        def _rec(oid, sku):
            return {"status": "REPLACE", "order_id": oid, "platform_order_id": "P-" + oid,
                    "shop": "袁州1(9352)", "shop_bcs": "袁州1", "paid_time": "2026-10-19 09:00",
                    "vc": "BCS-AAAA-1", "cn_name": "测试品", "matched_sku": "BCS-AAAA-1",
                    "expected_sku": sku, "note": ""}

        recs = [_rec("101", "SKU-1"), _rec("102", "SKU-2"), _rec("103", "SKU-3")]
        attempts = {"101": 0, "102": 0, "103": 0}

        def fake_items(cred, oid):
            attempts[oid] += 1
            if oid == "101" and attempts[oid] == 1:
                raise requests.exceptions.ReadTimeout("read timed out")
            if oid == "102":
                raise RuntimeError("orderItemId 数量=2，需人工处理")
            return ["item-" + oid]

        def fake_stock(cred, sku, wh, page_size=20):
            return None if sku == "SKU-3" else {"id": 900 + int(sku[-1]), "stockSku": sku}

        tmp = tempfile.mkdtemp()
        real_open = open
        try:
            def block_cn(path, *a, **kw):
                if "更换失败订单" in str(path):
                    raise PermissionError(13, "Permission denied", str(path))
                return real_open(path, *a, **kw)

            with mock.patch.object(mabang, "fetch_order_item_ids", side_effect=fake_items), \
                    mock.patch.object(mabang, "search_stock", side_effect=fake_stock), \
                    mock.patch.object(mabang, "replace_order_item") as mrep, \
                    mock.patch.object(mabang.time, "sleep", lambda *_: None):
                ok, fail = mabang.apply_replace({}, recs, 1457537)

            self.assertEqual((ok, fail), (1, 2))
            self.assertEqual(recs[0]["status"], "REPLACE", "重试成功应保持 REPLACE")
            self.assertIn("重试成功", recs[0]["note"])
            self.assertEqual(recs[1]["status"], "REPLACE_FAIL")
            self.assertIn("更换失败", recs[1]["note"])
            self.assertEqual(recs[2]["status"], "REPLACE_FAIL")
            self.assertIn("未找到库存", recs[2]["note"])
            self.assertEqual(attempts["101"], 2, "ReadTimeout 应重试一次")
            self.assertEqual(attempts["102"], 1, "确定性错误不应重试")
            self.assertEqual(mrep.call_count, 1, "只有重试成功那单真正提交了更换")

            # 失败专表：中文名被文件名策略拦截 → 回退 ASCII 别名，且只含 2 行失败明细
            with mock.patch.object(mabang.config, "LOG_DIR", tmp), \
                    mock.patch("builtins.open", side_effect=block_cn):
                path = mabang.report_replace_fail(recs)
            self.assertIn("mabang_replace_fail_", os.path.basename(path))
            with real_open(path, encoding="utf-8-sig") as f:
                rows = list(_csv.DictReader(f))
            self.assertEqual(len(rows), 2)
            self.assertTrue(all(r["失败原因"].startswith("更换失败") for r in rows))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_25l_mabang_process_chain_exclude_unit(self):
        """离线单测：① 失败单ID 传给②按ID排除；单独跑 forecast（无该属性）不受影响。"""
        import contextlib
        import io as _io
        from types import SimpleNamespace
        from unittest import mock
        from wb_ops.services.order import mabang, mabang_process

        def mk_order(oid):
            return {"id": oid, "shopIdText": "9352", "platformOrderId": "P-" + oid,
                    "paidTime": "2026-10-19 09:00:00"}

        def mk_rec(oid):
            return {"status": "REPLACE", "order_id": oid, "platform_order_id": "P-" + oid,
                    "shop": "袁州1(9352)", "shop_bcs": "袁州1", "paid_time": "2026-10-19 09:00",
                    "vc": "BCS-AAAA-1", "cn_name": "测试品", "matched_sku": "BCS-AAAA-1",
                    "expected_sku": "SKU-1", "note": ""}

        cred = {"shop_map": {"9352": "袁州1(9352)"}, "warehouse_id": 1457537}
        captured = {}

        def fake_apply(_cred, recs, _wh):
            recs[0]["status"] = "REPLACE_FAIL"
            recs[0]["note"] = "更换失败: Read timed out"
            return 0, 1

        args = SimpleNamespace(days=1, page_size=100, apply=True, wait=0,
                               upload_timeout=None, poll_interval=None,
                               url="http://feishu.example/base",
                               table="订单登记", no_pending=False)
        buf = _io.StringIO()
        with mock.patch.object(mabang, "_mabang_cred", return_value=cred), \
                mock.patch.object(mabang, "_api_ready", return_value=True), \
                mock.patch.object(mabang, "fetch_pending_orders",
                                  return_value=[mk_order("101")]), \
                mock.patch.object(mabang, "classify", return_value=[mk_rec("101")]), \
                mock.patch.object(mabang, "apply_replace", side_effect=fake_apply), \
                mock.patch.object(mabang, "report_no_sku"), \
                mock.patch.object(mabang, "report_replace_fail"), \
                mock.patch.object(mabang, "write_csv", return_value=""), \
                mock.patch.object(mabang, "run_forecast",
                                  side_effect=lambda ns: captured.update(ns=ns) or 0), \
                mock.patch.object(mabang_process.feishu_register, "run", return_value=0), \
                mock.patch.object(mabang_process.credentials, "get",
                                  return_value=SimpleNamespace(feishu_base_url=lambda: "http://x")), \
                contextlib.redirect_stdout(buf):
            code = mabang_process.run(args)
        out = buf.getvalue()
        self.assertEqual(code, 0)
        self.assertEqual(captured["ns"].replace_fail_order_ids, ["101"],
                         "步骤① 的失败单ID 应传给步骤②")
        self.assertIn("更换失败 1 单", out)
        self.assertIn("更换失败 1 单未预报/交运", out)

        def run_collect(with_ids):
            a = SimpleNamespace(days=1, page_size=100)
            if with_ids:
                a.replace_fail_order_ids = ["101"]
            b = _io.StringIO()
            with mock.patch.object(mabang, "_mabang_cred", return_value=cred), \
                    mock.patch.object(mabang, "fetch_pending_orders",
                                      return_value=[mk_order("101"), mk_order("102")]), \
                    mock.patch.object(mabang, "classify",
                                      return_value=[mk_rec("101"), mk_rec("102")]), \
                    mock.patch.object(mabang, "report_no_sku"), \
                    mock.patch("wb_ops.adapters.wb_client.fetch_canceled_ids",
                               return_value=set()), \
                    contextlib.redirect_stdout(b):
                target, _c = mabang.collect_forecast_orders(a)
            return [r["order_id"] for r in target], b.getvalue()

        ids_with, out_with = run_collect(True)
        self.assertEqual(ids_with, ["102"], "更换失败单应被排除出可预报")
        self.assertIn("更换失败=1", out_with)
        ids_alone, _ = run_collect(False)
        self.assertEqual(ids_alone, ["101", "102"],
                         "单独跑 mabang-forecast（无该属性）不应排除任何单")

    def test_25i_mabang_forecast_config_cached_unit(self):
        """离线单测：upload_forecast_batch 的 getForecastConfig 模板循环外只取一次。"""
        from unittest import mock
        from wb_ops.adapters import mabang_client as mc

        cfg = {"forecastLogistics": "2369||262534", "smt_sfc_shipping_method": "3"}
        posts = []

        class _Resp:
            def raise_for_status(self):
                pass

            def json(self):
                return {"success": True, "message": "批次已插入上传队列，请5-10分后查看预报结果"}

        def fake_post(url, params=None, headers=None, data=None, timeout=None):
            posts.append(dict(data or {}))
            return _Resp()

        with mock.patch.object(mc, "get_forecast_config", return_value=cfg) as mcfg, \
                mock.patch.object(mc, "aamz_headers", return_value={}), \
                mock.patch.object(mc.requests, "post", side_effect=fake_post), \
                mock.patch.object(mc.time, "sleep", lambda *_: None):
            msgs = mc.upload_forecast_batch({}, ["A", "B", "C"])

        self.assertEqual(mcfg.call_count, 1, "配置模板应只取一次（循环外）")
        self.assertEqual([p["batchNoInfo"] for p in posts], ["A", "B", "C"])
        self.assertTrue(all(p.get("wb_automark") == "1" and p.get("is_set_forecast") == "2"
                            for p in posts), "应保持自动发货参数")
        self.assertEqual(posts[0].get("forecastLogistics"), "2369||262534",
                         "应使用模板字段（未走降级路径）")
        self.assertTrue(msgs)

    def test_25f_mabang_process_find_batches_pagination_unit(self):
        """离线单测：find_batches 首批未命中时按 pageHtml 页数自动翻页并命中。"""
        from unittest import mock
        from wb_ops.services.order import mabang

        def side(cred, status=1, rows_per_page=200, page=1, **kw):
            if status != 3:
                return ([], {})
            if page == 1:
                return ([{"batchNo": f"X{i}"} for i in range(200)],
                        {"pageTotal": 201, "pageCount": 2, "pageCur": 1})
            return ([{"batchNo": "ZZ", "shopId": "主号7", "status": 3}],
                    {"pageTotal": 201, "pageCount": 2, "pageCur": 2})

        with mock.patch.object(mabang, "get_forecast_list", side_effect=side):
            res = mabang.find_batches(None, ["ZZ"], status=3, rows_per_page=200)

        self.assertTrue(res["all_found"], "翻页后应命中后段批次")
        self.assertEqual(res["pages_fetched"], 2)
        self.assertEqual(res["found"]["ZZ"]["shopId"], "主号7")

    # ---------------- 8. 客服与监控 ----------------
    def test_26_questions(self):
        res = self._run_cmd(["questions", "--no-detail", "--shops", "9352"], expect_code=0)
        self.assertIn("未处理提问", res.stdout)

    def test_27_questions_watch_once(self):
        res = self._run_cmd(["questions-watch", "--once", "--shops", "9352"], expect_code=0)
        self.assertIn("--once", res.stdout)

    def test_27b_appeals(self):
        res = self._run_cmd(["appeals", "--no-cn", "--shops", "9352", "--limit", "1"], expect_code=0)
        self.assertIn("店铺 1 个", res.stdout)
        self.assertIn("供应商代码", res.stdout)

    # ---------------- 9. 别名路由与并发安全 ----------------
    def test_28_cli_aliases(self):
        res1 = self._run_cmd(["mapping-merge", "--help"], expect_code=0)
        self.assertIn("merge", res1.stdout)
        res2 = self._run_cmd(["mapping-sync", "--help"], expect_code=0)
        self.assertIn("shops-mapping", res2.stdout)
        res3 = self._run_cmd(["order-pipeline", "--help"], expect_code=0)
        self.assertIn("mabang-process", res3.stdout)

    def test_29_mapping_rename(self):
        from wb_ops.storage.mapping_repo import MappingRepository
        try:
            res = self._run_cmd(["mapping-rename", "--vc", "BCS-TEST-TEST12345", "--cn", "单元测试商品-纠偏", "--reason", "自动化验证"], expect_code=0)
            self.assertIn("BCS-TEST-TEST12345", res.stdout)
            self.assertIn("单元测试商品-纠偏", res.stdout)
            # 2026-09-21：纠偏须同步写入全局已知池（vc_known.json），
            # 避免「前缀码匹配不到的怪 vc」其名字仅存于 vc_override.json 而随文件丢失失效
            self.assertIn("已知池同步", res.stdout)
            known = MappingRepository.load_vc_known()
            self.assertEqual((known.get("BCS-TEST-TEST12345") or {}).get("cn"), "单元测试商品-纠偏")
        finally:
            try:
                ov = MappingRepository.load_vc_override()
                if "BCS-TEST-TEST12345" in ov:
                    del ov["BCS-TEST-TEST12345"]
                    MappingRepository.save_vc_override(ov)
                kn = MappingRepository.load_vc_known()
                if "BCS-TEST-TEST12345" in kn:
                    del kn["BCS-TEST-TEST12345"]
                    MappingRepository.save_vc_known(kn)
            except Exception:
                pass

    def test_30_safe_io_concurrency(self):
        from wb_ops.framework.safe_io import FileLock, atomic_dump_json, safe_load_json
        test_path = os.path.join(BASE_DIR, "data", "state", ".test_ci_lock.json")
        try:
            with FileLock(test_path, timeout=5.0):
                # 测试重入锁与原子读写事务
                atomic_dump_json(test_path, {"test_run": True, "count": 1}, use_lock=True)
                d = safe_load_json(test_path, use_lock=True)
                self.assertEqual(d.get("count"), 1)
        finally:
            if os.path.exists(test_path):
                os.remove(test_path)


    # ---------------- 10. 补测：原缺测命令的只读路径 ----------------
    # 说明：本段全部为「只读/无副作用」路径 —— 参数校验、缺失文件报错、dry-run 预览、只读预览。
    #      绝不触发任何写平台/写飞书/写系统的动作（所有写操作命令均不带 --apply）。
    def test_31_import_shelve_argcheck(self):
        # 缺位置参数 → argparse 直接报错退出（不进入业务逻辑，不触发 ensure_snapshots/网络）
        res = self._run_cmd(["import-shelve"], expect_code=2)
        self.assertIn("usage", (res.stdout + res.stderr).lower())

    def test_32_discount_aliases_below(self):
        # 两个别名命令走 dry-run，并覆盖新增的双侧区间参数（--below 走 WB 升序接口）
        for alias in ("discount-wb", "discount-scan"):
            with self.subTest(alias=alias):
                res = self._run_cmd([alias, "--threshold", "55", "--below", "40", "--target", "50",
                                     "--shops", "9352", "--limit", "1"], expect_code=0)
                self.assertIn("dry-run", res.stdout)
                self.assertIn("<40", res.stdout)

    def test_33_ai_test_missing_qa(self):
        # 数据集文件不存在 → 明确报错退出（不调用 LLM，不产生费用）
        res = self._run_cmd(["ai-test", "--qa", os.path.join(BASE_DIR, "_scratch", "__no_such_qa__.json")],
                            expect_code=1)
        self.assertIn("[错误]", res.stdout + res.stderr)

    def test_34_feishu_register_dry_run(self):
        # scope=latest：默认合并「待处理订单」数据源（覆盖只匹配未进预报/上传/交运的单）
        res = self._run_cmd(["feishu-register", "--scope", "latest"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("[合并]", res.stdout)
        self.assertIn("真排除", res.stdout)

    def test_34b_feishu_register_no_pending(self):
        # --no-pending：关闭合并（只按 orderalllist 登记）
        res = self._run_cmd(["feishu-register", "--scope", "latest", "--no-pending"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertNotIn("[合并]", res.stdout)

    def test_34c_feishu_register_dedup_dry_run(self):
        # 存量清理命令默认 dry-run：全表分页读 + 打印统计；不传 --apply 不得删除任何记录
        res = self._run_cmd(["feishu-register-dedup"], expect_code=0)
        self.assertIn("[表格]", res.stdout)
        self.assertIn("[统计]", res.stdout)                    # 证明全表分页读通（防 2000 截断）
        self.assertTrue("[重复]" in res.stdout or "无重复记录" in res.stdout)
        self.assertNotIn("[写入] 开始删除", res.stdout)         # dry-run 绝不删除

    def test_34d_feishu_register_inventory_sku_unit(self):
        """离线单测：登记「库存SKU」取本地价格表（cn2sku），绝不取马帮系统匹配值。"""
        import tempfile
        from types import SimpleNamespace
        from unittest import mock
        from wb_ops.services.order import feishu_register

        order = {"id": "421000001", "shopIdText": "9352", "platformOrderId": "6060000001",
                 "paidTime": "2026-10-19 09:00:00",
                 "order_ellipsis_other_title": "商品编号*数量:BCS-AAAA-100000001*1",
                 "order_ellipsis_title": "商品编号*数量:BCS-WRNG-999*1"}   # 系统匹配是错的
        args = SimpleNamespace(scope="pending", days=1, page_size=100, no_pending=True,
                               date="", begin="", end="")
        tmp = tempfile.mkdtemp()
        with mock.patch.object(feishu_register.mabang, "_mabang_cred",
                               return_value={"shop_map": {"9352": "袁州1(9352)"}}), \
                mock.patch.object(feishu_register.mabang, "fetch_pending_orders",
                                  return_value=[order]), \
                mock.patch.object(feishu_register.mabang, "load_local_mapping",
                                  return_value=({"BCS-AAAA-100000001": "测试品"},
                                                {"测试品": "SKU-LOCAL-1"})), \
                mock.patch.object(feishu_register.config, "shop_json_path",
                                  return_value=os.path.join(tmp, "absent.json")), \
                mock.patch("wb_ops.adapters.wb_client.fetch_canceled_ids",
                           return_value=set()):
            records, _skip, _notice, _cred = feishu_register.build_records(args)

        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["库存SKU"], "SKU-LOCAL-1",
                         "库存SKU 必须来自本地价格表")
        self.assertNotEqual(records[0]["库存SKU"], "BCS-WRNG-999",
                            "不得回写马帮系统匹配值")

    def test_35_mabang_stock_register_dry_run(self):
        res = self._run_cmd(["mabang-stock-register"], expect_code=0)
        self.assertIn("dry-run", res.stdout)

    def test_36_mabang_stock_daily_end_only(self):
        # --end 单独使用必须报错中止（避免静默删列）
        res = self._run_cmd(["mabang-stock-daily", "--end", "2026-09-20"], expect_code=1)
        self.assertIn("--end", res.stdout)

    def test_36b_feishu_vc_stats_offline(self):
        """feishu-vc-stats 离线断言（**不联网、不调 lark-cli**）：字段解析 / 窗口 / 跨店合并 / 降序 / 前缀。

        背景：飞书「订单登记」表的 vendorCode 载体是 `BCS编号` 列（表内无「供应商代码」列）；
        同一 vendorCode 在各店 wb编号 不同，但统计须**跨店合并**到同一条（用户口径 2026-09-23）。
        """
        from datetime import datetime, timedelta

        from wb_ops.services.order import feishu_vc_stats as st

        # ① 字段解析：日期取前 10 位、店铺(list 多选)、订单量 number、BCS编号
        row = st.parse_row({
            "日期": "2026-09-22T03:26:00.000+08:00",
            "订单编号": "419915682",
            "店铺": ["袁州1"],
            "BCS编号": "BCS-JPAV-163906301",
            "商品中文名": "毛球修剪器",
            "订单量": 3,
        })
        self.assertEqual(row["日期"], "2026-09-22")
        self.assertEqual(row["vc"], "BCS-JPAV-163906301")
        self.assertEqual(row["cn"], "毛球修剪器")
        self.assertEqual(row["shops"], ["袁州1"])
        self.assertEqual(row["qty"], 3)
        self.assertEqual(row["order_id"], "419915682")
        # 缺字段 / 空店铺 容错
        empty = st.parse_row({})
        self.assertEqual((empty["日期"], empty["vc"], empty["shops"], empty["qty"]), ("", "", [], 1))

        # ② 窗口：闭区间 + 无日期不命中
        self.assertTrue(st.in_window({"日期": "2026-09-22"}, "2026-09-22", "2026-09-22"))
        self.assertTrue(st.in_window({"日期": "2026-09-17"}, "2026-09-17", "2026-09-23"))
        self.assertFalse(st.in_window({"日期": "2026-09-16"}, "2026-09-17", "2026-09-23"))
        self.assertFalse(st.in_window({"日期": ""}, "2026-09-17", "2026-09-23"))

        # ③ 窗口推导：days 含今天；date 单天；仅 begin 按单天；仅 end 报错
        b, e = st.resolve_window(days=7)
        self.assertEqual(e, time.strftime("%Y-%m-%d"))
        self.assertEqual(datetime.strptime(e, "%Y-%m-%d") - datetime.strptime(b, "%Y-%m-%d"),
                         timedelta(days=6))
        self.assertEqual(st.resolve_window(date="2026-9-5"), ("2026-09-05", "2026-09-05"))
        self.assertEqual(st.resolve_window(begin="2026-09-20"), ("2026-09-20", "2026-09-20"))
        self.assertEqual(st.resolve_window(begin="2026-09-02", end="2026-09-05"), ("2026-09-02", "2026-09-05"))
        with self.assertRaises(ValueError):
            st.resolve_window(end="2026-09-22")

        # ④ 跨店合并 + 单数降序 + 同数按 key 升序 + 件数累加
        rows = [
            {"日期": "2026-09-22", "vc": "BCS-AAAA-1", "cn": "甲", "shops": ["袁州1"], "qty": 2, "order_id": "o1"},
            {"日期": "2026-09-22", "vc": "BCS-AAAA-1", "cn": "甲", "shops": ["袁州3"], "qty": 1, "order_id": "o2"},
            {"日期": "2026-09-22", "vc": "BCS-BBBB-2", "cn": "乙", "shops": ["袁州2"], "qty": 1, "order_id": "o3"},
            {"日期": "2026-09-22", "vc": "BCS-CCCC-3", "cn": "丙", "shops": ["袁州1"], "qty": 1, "order_id": "o4"},
            {"日期": "2026-09-22", "vc": "", "cn": "", "shops": ["袁州1"], "qty": 1, "order_id": "o5"},
        ]
        agg = st.aggregate(rows)
        self.assertEqual([(r["key"], r["orders"], r["qty"]) for r in agg],
                         [("BCS-AAAA-1", 2, 3), ("BCS-BBBB-2", 1, 1), ("BCS-CCCC-3", 1, 1), (st.NO_VC_KEY, 1, 1)])
        self.assertEqual([r["rank"] for r in agg], [1, 2, 3, 4])
        # 未归类桶恒排最后（不然 '(' 码位小于 'B' 会插到真实 vendorCode 前面）
        self.assertTrue(str(agg[-1]["key"]).startswith("("))
        # 跨店合并：同一 vc 在袁州1/袁州3 的店铺分布已合并（各 1 单 → 按单数降序、同数按名升序）
        self.assertEqual(agg[0]["shops"], ["袁州1", "袁州3"])
        self.assertEqual(agg[0]["cn"], "甲")
        # 无 BCS编号 的记录不静默丢弃，单独成桶
        self.assertEqual([r["key"] for r in agg if r["key"] == st.NO_VC_KEY], [st.NO_VC_KEY])

        # ⑤ 前缀码聚合视角
        pre = st.aggregate(rows, by_prefix=True)
        self.assertEqual([(r["key"], r["orders"]) for r in pre], [("AAAA", 2), ("BBBB", 1), ("CCCC", 1), (st.NO_PREFIX_KEY, 1)])

        # ⑥ 店铺过滤：空集 = 不过滤；否则取交集
        r_shop = {"日期": "2026-09-22", "vc": "x", "shops": ["袁州2"]}
        self.assertTrue(st.match_shops(r_shop, None))
        self.assertTrue(st.match_shops(r_shop, set()))
        self.assertFalse(st.match_shops(r_shop, {"袁州1", "袁州3"}))
        self.assertTrue(st.match_shops(r_shop, {"袁州2"}))

        # ⑦ 店铺标签解析：9352 → 袁州1（数字经 credentials 反查）；未知数字保留原值
        self.assertEqual(st.resolve_shop_labels("9352"), {"袁州1"})
        self.assertEqual(st.resolve_shop_labels("袁州1,袁州3"), {"袁州1", "袁州3"})
        self.assertEqual(st.resolve_shop_labels(["9353", "璧山1"]), {"袁州2", "璧山1"})
        self.assertEqual(st.resolve_shop_labels(""), set())

    def test_37_cookies_update_missing_file(self):
        res = self._run_cmd(["cookies-update", "__no_such_cookies__.md"], expect_code=1)
        self.assertIn("[错误]", res.stdout + res.stderr)

    def test_37b_cookie_refresh_wb_dry_run(self):
        res = self._run_cmd(["cookie-refresh-wb", "--shops", "999999999"], expect_code=0)
        self.assertIn("[汇总] 成功 0/0", res.stdout)

    def test_38_daily_safe_scope(self):
        # 用不存在的店铺 ID 限定，使各步骤「无匹配店铺」空转（不产生任何写操作）；
        # daily 的输出落当日日志 data/logs/daily_YYYYMMDD.log，故断言日志内容
        res = self._run_cmd(["daily", "check", "--shops", "0"], expect_code=0)
        self.assertEqual(res.stdout.strip(), "")  # 步骤输出被重定向进日志
        log_path = os.path.join(BASE_DIR, "data", "logs", "daily_" + time.strftime("%Y%m%d") + ".log")
        self.assertTrue(os.path.exists(log_path), f"未生成当日日志: {log_path}")
        with open(log_path, encoding="utf-8", errors="replace") as f:
            blob = f.read()
        self.assertIn("[改折扣]", blob)
        self.assertIn("店铺 0 个", blob)

    def test_39_schedule_plan(self):
        res = self._run_cmd(["schedule", "--plan"], expect_code=0)
        self.assertIn("[dry-run]", res.stdout)
        self.assertIn("WB_Daily_Morning", res.stdout)
        self.assertIn("schtasks", res.stdout)

    def test_40_shelve_dry_run(self):
        # 支持单品/多品、指定前缀码与尺寸重量等灵活参数
        res = self._run_cmd(["shelve", "248364237", "388854754", "--price", "59", "--prefix", "ABCD", "--dims", "10*20*30/0.5"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("248364237", res.stdout)
        self.assertIn("388854754", res.stdout)

    def test_41_shelve_old_dry_run(self):
        # 支持指定任意自定义完整 vendorCode
        res = self._run_cmd(["shelve-old", "248364237", "--price", "59", "--vc", "BCS-CUSTOM-SPECIAL-12345", "--dims", "10*20*30/0.5"], expect_code=0)
        self.assertIn("dry-run", res.stdout)
        self.assertIn("BCS-CUSTOM-SPECIAL-12345", res.stdout)

    def test_41b_shelve_old_basket_path_unit(self):
        """离线单测：basket 分片表对齐官方插件 65 区间；vol 越界回退表尾分片（不再 768 外推）。

        背景（2026-10-19）：旧表只到 `24965→60` 且越界用 768 步长外推，vol>24965 会算出
        不存在的 61/62… 分片（card.json / 图片全 404）。官方 1.3.4 扩到 65 区间且越界回退表尾。
        """
        from wb_ops.adapters import wb_client as wc

        self.assertEqual(len(wc._BASKET_TABLE), 65)
        self.assertEqual(wc._BASKET_TABLE[-1], (28805, "65"))
        cases = {
            1931583426: "https://basket-53.wbbasket.ru/vol19315/part1931583/1931583426",
            1457749715: "https://basket-47.wbbasket.ru/vol14577/part1457749/1457749715",
            1939999999: "https://basket-53.wbbasket.ru/vol19399/part1939999/1939999999",
            2900000000: "https://basket-65.wbbasket.ru/vol29000/part2900000/2900000000",
        }
        for nm, want in cases.items():
            self.assertEqual(wc.basket_base(nm), want, f"nm={nm}")

    def test_42_shelve_backend_interchangeable(self):
        # 验证 replicate 与 import_shelve 可平滑无缝将 shelve_backend 切换为 shelve_old
        from wb_ops.services.replicate import replicate, import_shelve
        from wb_ops.services.replicate import shelve_new, shelve_old
        self.assertTrue(hasattr(shelve_new, "execute_shelve"))
        self.assertTrue(hasattr(shelve_old, "execute_shelve"))
        # 验证两者签名与关键参数完全对称
        import inspect
        sig_new = inspect.signature(shelve_new.execute_shelve)
        sig_old = inspect.signature(shelve_old.execute_shelve)
        self.assertEqual(list(sig_new.parameters.keys()), list(sig_old.parameters.keys()))

    def test_42b_shelve_file_dp_floor_unit(self):
        """shelve --file：CSV 上架价按商品价格表「实际售价」floor（离线）。

        口径 2026-10-20：D 列改为实际售价后，`shelve --file` 与 `price`/主链路口径统一为
        floor(实际售价)——此前该分支直接 float(dp) 不取整（88.9 → 88.9，现 → 88）。
        """
        import os
        import shutil
        import tempfile
        from unittest import mock

        import openpyxl
        from wb_ops import config
        from wb_ops.services.replicate import shelve_common

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["产品图片", "卖家SKU", "产品中文名", "实际售价", "尺寸",
                   "最低售价", "vendorCode前缀码", "库存SKU"])
        ws.append(["", "SKU-9", "测试品", 88.9, "10*10*10/0.1", 40, "TEST", "s9"])

        tmp = tempfile.mkdtemp()
        try:
            boss_path = os.path.join(tmp, "boss.xlsx")
            wb.save(boss_path)
            csv_path = os.path.join(tmp, "orders.csv")
            with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
                f.write("商品中文名,现在商品的nmId,VC\n测试品,123456,OTHER-1\n")

            with mock.patch("openpyxl.load_workbook", return_value=wb), \
                    mock.patch.object(config, "BOSS_XLSX", boss_path):
                items, _raw = shelve_common.load_shelve_file(csv_path)

            self.assertEqual(len(items), 1)
            self.assertEqual(items[0]["price"], 88)
            self.assertIsInstance(items[0]["price"], int)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_56a_merge_dp_header_new_unit(self):
        """merge 读映射总表：价格列按「实际售价」表头定位（挪到非 C 列也能读到）。"""
        import io as _io
        import os
        import shutil
        import tempfile

        import openpyxl
        from unittest import mock
        from wb_ops import config
        from wb_ops.framework import xlsx_headers
        from wb_ops.storage.mapping_repo import MappingRepository

        xlsx_headers.reset_warnings()
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "mapping.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "映射总表"
            # 实际售价 故意放在第 5 列（非 C 列），验证不依赖位置
            ws.append(["产品中文名", "vendorCode", "店铺9352价格(CNY)", "折扣%", "实际售价"])
            ws.append(["测试品", "BCS-AAAA-1", 100, 10, 88.9])
            wb.save(path)

            with mock.patch.object(config, "MAPPING_XLSX", path), \
                    mock.patch.object(MappingRepository, "load_vc_override", return_value={}):
                state, _ex = MappingRepository.load_mapping_state()

            self.assertEqual(state["BCS-AAAA-1"]["dp"], 88.9)
            self.assertEqual(state["BCS-AAAA-1"]["shop_price"], 100)
            self.assertEqual(state["BCS-AAAA-1"]["cn"], "测试品")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_56b_merge_dp_header_legacy_warn_unit(self):
        """merge 读映射总表：只有旧列名「双倍售价」时不崩、dp=None，并给出可见警告。"""
        import contextlib
        import io as _io
        import os
        import shutil
        import tempfile

        import openpyxl
        from unittest import mock
        from wb_ops import config
        from wb_ops.framework import xlsx_headers
        from wb_ops.storage.mapping_repo import MappingRepository

        xlsx_headers.reset_warnings()
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "mapping_old.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "映射总表"
            ws.append(["产品中文名", "vendorCode", "双倍售价", "店铺9352价格(CNY)"])
            ws.append(["测试品", "BCS-AAAA-2", 88.9, 100])
            wb.save(path)

            err = _io.StringIO()
            with mock.patch.object(config, "MAPPING_XLSX", path), \
                    mock.patch.object(MappingRepository, "load_vc_override", return_value={}), \
                    contextlib.redirect_stderr(err):
                state, _ex = MappingRepository.load_mapping_state()

            self.assertIsNone(state["BCS-AAAA-2"]["dp"])       # 只认新名
            self.assertEqual(state["BCS-AAAA-2"]["shop_price"], 100)
            msg = err.getvalue()
            self.assertIn("实际售价", msg)
            self.assertIn("双倍售价", msg)
            self.assertIn("merge", msg)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_56c_merge_single_shop_dp_header_unit(self):
        """merge 读单店表：价格列按「实际售价」表头定位；旧列名 → dp=None + 警告。"""
        import contextlib
        import io as _io
        import os
        import shutil
        import tempfile

        import openpyxl
        from unittest import mock
        from wb_ops import config
        from wb_ops.framework import xlsx_headers
        from wb_ops.services.catalog import mapping_excel

        tmp = tempfile.mkdtemp()
        try:
            def build(name, dp_header):
                p = os.path.join(tmp, name)
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "在架映射明细"
                # 实际售价 放第 4 列（非第 3 列）
                ws.append(["产品中文名", "vendorCode", "本店售价(CNY)", dp_header, "本店库存"])
                ws.append(["测试品", "BCS-BBBB-1", 120, 88.9, 5])
                wb.save(p)
                return p

            p_new = build("new.xlsx", "实际售价")
            xlsx_headers.reset_warnings()
            with mock.patch.object(config, "shop_mapping_xlsx",
                                   lambda sid, name=None: p_new):
                rows, _unmapped = mapping_excel.load_single_shop_rows(9352)
            self.assertEqual(rows["BCS-BBBB-1"]["dp"], 88.9)
            self.assertEqual(rows["BCS-BBBB-1"]["price"], 120)

            p_old = build("old.xlsx", "双倍售价")
            xlsx_headers.reset_warnings()
            err = _io.StringIO()
            with mock.patch.object(config, "shop_mapping_xlsx",
                                   lambda sid, name=None: p_old), \
                    contextlib.redirect_stderr(err):
                rows_old, _u = mapping_excel.load_single_shop_rows(9352)
            self.assertIsNone(rows_old["BCS-BBBB-1"]["dp"])
            self.assertIn("实际售价", err.getvalue())
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_57_importshelve_foreign_legacy_dp_raise_unit(self):
        """import-shelve 读他人表：旧列名「双倍售价」→ 直接报错；新名挪位仍可读。"""
        import os
        import shutil
        import tempfile

        import openpyxl
        from wb_ops.framework import xlsx_headers
        from wb_ops.services.replicate import foreign_table
        from wb_ops.framework.xlsx_headers import HeaderError

        tmp = tempfile.mkdtemp()
        try:
            def build(name, dp_header):
                p = os.path.join(tmp, name)
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "映射总表"
                ws.append(["产品中文名", "vendorCode", dp_header, "WB商品码"])
                ws.append(["他人品", "BCS-CCCC-1396678", 66.6, 1396678])
                wb.save(p)
                return p

            xlsx_headers.reset_warnings()
            with self.assertRaises((RuntimeError, HeaderError)) as ctx:
                foreign_table.load_foreign(build("old.xlsx", "双倍售价"))
            msg = str(ctx.exception)
            self.assertIn("实际售价", msg)
            self.assertIn("双倍售价", msg)
            self.assertIn("实际表头", msg)

            items, _bad, _no_nm = foreign_table.load_foreign(build("new.xlsx", "实际售价"))
            self.assertEqual(len(items), 1)
            self.assertEqual(items[0]["dp"], 66.6)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_58_price_boss_dp_column_reorder_unit(self):
        """price 读商品价格表：价格列按「实际售价」表头定位，四个读取点都位移仍正确。"""
        import os
        import shutil
        import tempfile

        import openpyxl
        from unittest import mock
        from wb_ops import config
        from wb_ops.framework import xlsx_headers
        from wb_ops.framework.xlsx_headers import HeaderError
        from wb_ops.services.replicate import shelve_common
        from wb_ops.storage.mapping_repo import MappingRepository

        tmp = tempfile.mkdtemp()
        try:
            def build(name, headers, row):
                p = os.path.join(tmp, name)
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Sheet1"
                ws.append(headers)
                ws.append(row)
                wb.save(p)
                return p

            headers = ["产品图片", "卖家SKU", "产品中文名", "尺寸", "最低售价",
                       "实际售价", "vendorCode前缀码", "库存SKU"]
            row = ["", "SKU-9", "测试品", "10*10*10/0.1", 40, 88.9, "TEST", "s9"]
            p_ok = build("boss_ok.xlsx", headers, row)

            xlsx_headers.reset_warnings()
            with mock.patch.object(config, "BOSS_XLSX", p_ok):
                boss = MappingRepository.load_boss()
                pmap = MappingRepository.load_prefix_map()
                details = MappingRepository.load_boss_details()
                shelve_dict = shelve_common.load_boss_shelve_dict()

            self.assertEqual(boss[0]["dp"], 88.9)
            self.assertEqual(boss[0]["floor"], 88)
            self.assertEqual(boss[0]["stock_sku"], "s9")
            self.assertEqual(pmap["TEST"]["dp"], 88.9)
            self.assertEqual(details["by_cn"]["测试品"]["dp"], 88.9)
            self.assertEqual(details["by_cn"]["测试品"]["price"], 40)   # 最低售价按列名取
            self.assertEqual(shelve_dict["测试品"]["dp"], 88.9)

            # 缺「实际售价」列 → 必须抛错（且不被 load_prefix_map 的宽 except 吞掉）
            headers_bad = ["产品图片", "卖家SKU", "产品中文名", "尺寸", "最低售价",
                           "vendorCode前缀码", "库存SKU"]
            row_bad = ["", "SKU-9", "测试品", "10*10*10/0.1", 40, "TEST", "s9"]
            p_bad = build("boss_bad.xlsx", headers_bad, row_bad)
            with mock.patch.object(config, "BOSS_XLSX", p_bad):
                with self.assertRaises(HeaderError):
                    MappingRepository.load_boss()
                with self.assertRaises(HeaderError):
                    MappingRepository.load_prefix_map()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_59_feishu_register_cn_fallback_unit(self):
        """feishu-register 填「商品中文名」三级回退：映射总表 → vc_known → 四位前缀码（离线）。

        2026-10-20：商品下架后 vc 从映射总表消失，登记中文名变空（74 单黑丝假发）。
        登记侧改为三级回退，仅为填名；马帮匹配/预报仍要求 vc 在映射表（在架）。
        """
        import os
        import shutil
        import tempfile

        import openpyxl
        from unittest import mock
        from wb_ops import config
        from wb_ops.services.order import feishu_register as fr
        from wb_ops.storage.mapping_repo import MappingRepository

        tmp = tempfile.mkdtemp()
        try:
            boss = os.path.join(tmp, "boss.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(["产品图片", "卖家SKU", "产品中文名", "实际售价", "尺寸",
                       "最低售价", "vendorCode前缀码", "库存SKU"])
            ws.append(["", "SZLK-1", "手足裂口霜", 80, "10*10*10/0.1", 40, "WENR", "szlk"])
            wb.save(boss)

            known = {"BCS-BBBB-2": {"cn": "归属池品", "dp": 88}}
            with mock.patch.object(config, "BOSS_XLSX", boss), \
                    mock.patch.object(MappingRepository, "load_vc_known", return_value=known):
                resolve = fr._cn_resolver({"BCS-AAAA-1": "映射表品"})
                self.assertEqual(resolve("BCS-AAAA-1"), "映射表品")        # ① 映射总表优先
                self.assertEqual(resolve("BCS-BBBB-2"), "归属池品")        # ② vc_known 兜底
                self.assertEqual(resolve("BCS-WENR-999"), "手足裂口霜")    # ③ 前缀码兜底
                self.assertIsNone(resolve("BCS-ZZZZ-123456"))              # 三级皆无
                self.assertIsNone(resolve(""))
                self.assertIsNone(resolve(None))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_43_legacy_payload_images_subject_id(self):
        # 验证旧接口 build_legacy_payload 在无外部图片与合成 detail 时，仍能保证图片非空且 subjectId 提取正确
        from wb_ops.domain.models import ShelveItem
        from wb_ops.services.replicate.shelve_old import build_legacy_payload
        from wb_ops.services.replicate.shelve_common import fetch_card_json
        item = ShelveItem(
            nm_id=248364237,
            price=99.0,
            vendor_code="BCS-CUSTOM-PAYLOAD-TEST",
            length=10, width=20, height=30, weight=0.5
        )
        card = fetch_card_json(248364237) or {"imt_name": "测试商品", "imt_id": 1, "data": {"subject_id": 2290}}
        payload, err = build_legacy_payload(item, [9352], {9352: 1929635}, card, None)
        self.assertIsNone(err)
        self.assertIsNotNone(payload)
        sdata = payload["shopDatas"][0]
        self.assertTrue(len(sdata["images"]) > 0, "图片链接不得为空")
        self.assertTrue(len(sdata["mainImage"]) > 0, "主图链接不得为空")
        self.assertGreater(sdata["subjectId"], 0, "类目ID必须大于0")
        self.assertEqual(sdata["vendorCode"], "BCS-CUSTOM-PAYLOAD-TEST")

    def test_44_legacy_payload_missing_warehouse(self):
        # 验证仓库配置缺失时返回明确错误而非崩溃抛出 KeyError
        from wb_ops.domain.models import ShelveItem
        from wb_ops.services.replicate.shelve_old import build_legacy_payload
        item = ShelveItem(nm_id=248364237, price=99.0, vendor_code="BCS-TEST-VC")
        card = {"data": {"subject_id": 100}}
        payload, err = build_legacy_payload(item, [99999], {}, card, None)
        self.assertIsNone(payload)
        self.assertIn("缺少对应发货仓库配置", err)

    def test_45_shelve_old_csv_file_support(self):
        # 验证 shelve-old 直接读取订单筛选 CSV 文件并进行预览
        csv_file = r"C:\Users\Admin\Desktop\筛选非我店订单_20260917.csv"
        if os.path.exists(csv_file):
            res = self._run_cmd(["shelve-old", "--file", csv_file, "--shops", "9352,9353,9356"], expect_code=0)
            self.assertIn("dry-run", res.stdout)
            self.assertIn("1378756362", res.stdout)
            self.assertIn("BCS-DZJF-209792332", res.stdout)
            self.assertIn("原始VC", res.stdout)

    def test_46_promo_gap_offline(self):
        """promo-gap 离线差集/分组断言（**不联网、不调 lark-cli**）。

        口径（2026-09-23 与用户确认并**修正为「合计」**）：**销量判据 = 该供应商代码在目标店铺
        （默认袁州1/2/3）的 7 天单数「合计」**，不是单店单数（同一 vc 各店 wb编号 不同，但它是同一个商品）。
        正向 gap：合计 ≥ 阈值 且某目标店未推广 → 该店该补推；
        反向 waste：合计 < 阈值 → 这些店所有「在投(status=9)」活动里的推广都建议关闭；
        nm 本地真源反查不到 vc → 「无法判定」（不推断）；暂停(11)/未识别状态 → 仅计数不逐条列。
        """
        from wb_ops.services.discount import promo_gap as pg

        # ① 阈值解析
        self.assertEqual(pg.resolve_min(None), 4)
        self.assertEqual(pg.resolve_min(""), 4)
        self.assertEqual(pg.resolve_min("6"), 6)
        for bad in (0, -3, "abc"):
            with self.assertRaises(ValueError):
                pg.resolve_min(bad)

        # ② 店铺选择：店铺ID 与短名双认
        all_shops = [{"shopId": 9352, "shopName": "袁州1"}, {"shopId": 9353, "shopName": "袁州2"}]
        self.assertEqual([x[1] for x in pg.select_shops(all_shops, "9352")], [9352])
        self.assertEqual([x[2] for x in pg.select_shops(all_shops, "袁州2")], ["袁州2"])
        self.assertEqual(len(pg.select_shops(all_shops, "9352,袁州2")), 2)
        self.assertEqual(len(pg.select_shops(all_shops, "")), 2)

        rows = [
            # AAAA：袁州1×2 + 袁州2×2 → 合计 4（单店各 2 均 <4 ⇒ 旧「单店」口径会漏判）
            {"日期": "2026-09-22", "vc": "BCS-AAAA-1", "cn": "甲", "shops": ["袁州1"], "qty": 1, "order_id": "a1"},
            {"日期": "2026-09-22", "vc": "BCS-AAAA-1", "cn": "甲", "shops": ["袁州1"], "qty": 2, "order_id": "a2"},
            {"日期": "2026-09-22", "vc": "BCS-AAAA-1", "cn": "甲", "shops": ["袁州2"], "qty": 1, "order_id": "a3"},
            {"日期": "2026-09-22", "vc": "BCS-AAAA-1", "cn": "甲", "shops": ["袁州2"], "qty": 1, "order_id": "a4"},
            # BBBB：袁州1×4 → 合计 4
            {"日期": "2026-09-22", "vc": "BCS-BBBB-2", "cn": "乙", "shops": ["袁州1"], "qty": 1, "order_id": "b1"},
            {"日期": "2026-09-22", "vc": "BCS-BBBB-2", "cn": "乙", "shops": ["袁州1"], "qty": 1, "order_id": "b2"},
            {"日期": "2026-09-22", "vc": "BCS-BBBB-2", "cn": "乙", "shops": ["袁州1"], "qty": 1, "order_id": "b3"},
            {"日期": "2026-09-22", "vc": "BCS-BBBB-2", "cn": "乙", "shops": ["袁州1"], "qty": 1, "order_id": "b4"},
            # CCCC：袁州1×1 → 合计 1
            {"日期": "2026-09-22", "vc": "BCS-CCCC-3", "cn": "丙", "shops": ["袁州1"], "qty": 1, "order_id": "c1"},
            # EEEE：店铺多选 袁州1+袁州3 ×4 → 合计只计一次=4（不是 8），拆分各计 4
            {"日期": "2026-09-22", "vc": "BCS-EEEE-4", "cn": "丁", "shops": ["袁州1", "袁州3"], "qty": 1, "order_id": "e1"},
            {"日期": "2026-09-22", "vc": "BCS-EEEE-4", "cn": "丁", "shops": ["袁州1", "袁州3"], "qty": 1, "order_id": "e2"},
            {"日期": "2026-09-22", "vc": "BCS-EEEE-4", "cn": "丁", "shops": ["袁州1", "袁州3"], "qty": 1, "order_id": "e3"},
            {"日期": "2026-09-22", "vc": "BCS-EEEE-4", "cn": "丁", "shops": ["袁州1", "袁州3"], "qty": 1, "order_id": "e4"},
            # FFFF：只在非目标店（璧山1）→ 目标店口径下不计
            {"日期": "2026-09-22", "vc": "BCS-FFFF-6", "cn": "己", "shops": ["璧山1"], "qty": 1, "order_id": "f1"},
            # 无 BCS编号
            {"日期": "2026-09-22", "vc": "", "cn": "", "shops": ["袁州1"], "qty": 1, "order_id": "n1"},
        ]

        # ③ 跨店合计 vs 各店拆分（★ 本次口径修正的核心）
        by_vc, by_shop, no_vc = pg.aggregate_orders(rows, {"袁州1", "袁州2", "袁州3"})
        self.assertEqual(by_vc["BCS-AAAA-1"]["orders"], 4)
        self.assertEqual(by_vc["BCS-AAAA-1"]["qty"], 5)
        self.assertEqual(by_shop["BCS-AAAA-1"]["袁州1"]["orders"], 2)
        self.assertEqual(by_shop["BCS-AAAA-1"]["袁州2"]["orders"], 2)
        self.assertEqual(by_vc["BCS-EEEE-4"]["orders"], 4)          # 多店铺记录合计只计一次
        self.assertEqual(by_shop["BCS-EEEE-4"]["袁州1"]["orders"], 4)
        self.assertEqual(by_shop["BCS-EEEE-4"]["袁州3"]["orders"], 4)
        self.assertNotIn("BCS-FFFF-6", by_vc)                        # 非目标店不计
        self.assertEqual(no_vc, 1)
        self.assertEqual(pg.shop_split_text(by_shop["BCS-AAAA-1"]), "袁州1=2,袁州2=2")
        by1, _ = pg.shop_rows_by_vc(rows, "袁州1")                    # 单店视角包装仍可用
        self.assertEqual(by1["BCS-BBBB-2"]["orders"], 4)

        # ④ 正向 gap：合计≥4 且该目标店未推广 → 该店缺口
        targets2 = [({}, 9352, "袁州1"), ({}, 9353, "袁州2")]
        adverted = {9352: {"BCS-BBBB-2"}, 9353: set()}
        snaps = {9352: {"BCS-AAAA-1": {}, "BCS-EEEE-4": {"nmId": 444}},
                 9353: {"BCS-AAAA-1": {"nmID": 111}}}
        g = pg.build_gap(targets2, rows, adverted, snaps, min_orders=4)
        self.assertEqual([(x["vc"], x["shop_id"]) for x in g["gap_listed"]],
                         [("BCS-AAAA-1", 9352), ("BCS-AAAA-1", 9353), ("BCS-EEEE-4", 9352)])
        self.assertEqual([(x["vc"], x["shop_id"]) for x in g["gap_unlisted"]],
                         [("BCS-BBBB-2", 9353), ("BCS-EEEE-4", 9353)])
        self.assertEqual(g["gap_listed"][0]["orders"], 4)            # 合计单数
        self.assertEqual(g["gap_listed"][0]["shop_orders"], 2)       # 该店单数
        self.assertEqual(g["gap_listed"][0]["shop_split"], "袁州1=2,袁州2=2")
        self.assertEqual(g["gap_listed"][1]["nm_id"], 111)           # nmID 键兼容
        self.assertEqual(g["gap_listed"][2]["nm_id"], 444)
        self.assertIn("nmId", g["gap_listed"][0]["备注"])             # 在架但缺 nmId
        s = g["gap_stats"]
        self.assertEqual((s["候选vc数"], s["已推广店位"], s["候选店铺位"]), (3, 1, 5))
        self.assertEqual((s["可直接补推"], s["需先上架"], s["缺nmId"]), (3, 2, 1))
        # 快照文件缺失 → 全部归组②并标注（不误判为「可补推」）
        g2 = pg.build_gap(targets2, rows, {9352: set(), 9353: set()},
                          {9352: None, 9353: None}, min_orders=4)
        self.assertEqual(g2["gap_listed"], [])
        self.assertIn("快照", g2["gap_unlisted"][0]["备注"])
        # 阈值降到 1：CCCC（合计 1）也成候选
        g3 = pg.build_gap(targets2, rows, {9352: set(), 9353: set()}, snaps, min_orders=1)
        self.assertIn("BCS-CCCC-3", {x["vc"] for x in g3["gap_listed"] + g3["gap_unlisted"]})

        # ⑤ 反向 waste：合计 ≥ 阈值 → 不关；合计 < 阈值 → 在投活动里的推广建议关闭
        targets3 = [({}, 9352, "袁州1"), ({}, 9353, "袁州2"), ({}, 9356, "袁州3")]

        def _prod(sid, sname, cid, status, nm, vc):
            return {"shop_id": sid, "shop_name": sname, "campaign_id": cid,
                    "campaign_name": "c%d" % cid, "status_id": status, "payment_model": "cpc",
                    "budget": 100, "create_date": "", "nm": nm, "vc": vc, "cn": "名%d" % nm,
                    "src": "店快照", "ru": "R%d" % nm, "subject": "", "fbo": 0, "mp": 0}

        prods = [
            _prod(9352, "袁州1", 1, 9, 11, "BCS-AAAA-1"),     # 合计 4 → 不关
            _prod(9352, "袁州1", 1, 9, 12, "BCS-BBBB-2"),     # 合计 4 → 不关
            _prod(9352, "袁州1", 2, 11, 13, "BCS-CCCC-3"),    # 合计 1 但已暂停 → 仅计数
            _prod(9352, "袁州1", 3, 9, 14, ""),               # vc 反查不到 → 无法判定
            _prod(9356, "袁州3", 4, 9, 15, "BCS-CCCC-3"),     # 合计 1 + 在投 → 建议关闭
            _prod(9356, "袁州3", 5, 9, 15, "BCS-CCCC-3"),     # 同 nm 的第二个活动
            _prod(9352, "袁州1", 6, 4, 16, "BCS-CCCC-3"),     # 未识别状态 → 已暂停 + 告警
        ]
        w = pg.build_waste(targets3, rows, prods, min_orders=4)
        self.assertEqual([(x["shop_id"], x["campaign_id"]) for x in w["waste_close"]],
                         [(9356, 4), (9356, 5)])
        self.assertTrue(all(x["verdict"] == pg.VERDICT_CLOSE for x in w["waste_close"]))
        self.assertEqual(w["waste_close"][0]["orders"], 1)            # 三店合计单数
        self.assertEqual(w["waste_close"][0]["shop_orders"], 0)       # 袁州3 本店 0 单
        self.assertEqual(w["waste_close"][0]["shop_split"], "袁州1=1")
        self.assertIn("低于阈值", w["waste_close"][0]["备注"])
        self.assertNotIn(1, [x["campaign_id"] for x in w["waste_close"]])   # 合计达标不进清单
        self.assertEqual([x["campaign_id"] for x in w["waste_paused"]], [2, 6])
        self.assertEqual(len(w["waste_unknown"]), 1)
        self.assertEqual(w["waste_unknown_status"], [(4, 9352, 6)])
        self.assertEqual(w["waste_stats"]["被推广商品位"], 7)
        self.assertEqual(w["waste_stats"]["建议关闭"], 2)

        # ⑥ 商品级汇总：同 (店铺, nm) 跨活动合并
        bn = w["waste_by_nm"]
        self.assertEqual([(x["shop_id"], x["nm"], x["n_campaigns"], x["orders"]) for x in bn],
                         [(9356, 15, 2, 1)])

        # ⑦ 活动级汇总：建议关闭数 == 被推广商品数 → 整活动建议关闭
        bc = {(x["shop_id"], x["campaign_id"]): x for x in w["waste_by_campaign"]}
        self.assertEqual(bc[(9356, 4)]["被推广商品数"], 1)
        self.assertEqual(bc[(9356, 4)]["建议关闭数"], 1)
        self.assertTrue(bc[(9356, 4)]["是否整活动建议关闭"])
        self.assertEqual(bc[(9352, 1)]["建议关闭数"], 0)              # 合计达标 → 不关
        self.assertEqual(bc[(9352, 1)]["是否整活动建议关闭"], "")
        self.assertEqual(bc[(9352, 2)]["建议关闭数"], 0)              # 暂停 → 不算建议关闭

    def test_36c_fbs_delivery_offline(self):
        """fbs-delivery 离线断言（**不联网**）：清单/商品解析、状态过滤、时间窗口、分页防死循环、聚合。

        数据形态取自 HAR 抓包样本（wb订单列表配送中.har / wb订单列表配送中2.har）：
        商品级状态实测分布 readyForPickup×3 / cancel×3 / shipment×1 / sold×1。
        口径（2026-10-08 与用户确认）：默认只统计 `shipment`（「待发货」）；时间基准 = 清单 createdAt；
        `orders/delivery` 返回该清单下全部订单（含已取消），必须按 state.code 过滤。
        """
        from wb_ops.adapters import wb_supply_client as sup
        from wb_ops.services.order import fbs_delivery as fd

        # ① 清单解析（supplyID / createdAt / state / 件数）
        ps = fd.parse_supply({
            "supplyID": "WB-GI-287405573", "name": "Заказ",
            "createdAt": "2026-10-03T10:42:05Z",
            "state": {"code": "onway", "title": "前往接待处途中"},
            "ordersCnt": 8, "parcelsCnt": 1, "boxesCnt": 1,
            "warehouse": "莫斯科仓库",
            "destinationWarehouse": {"name": "Москва"},
        })
        self.assertEqual(ps["supply_id"], "WB-GI-287405573")
        self.assertEqual(ps["created_day"], "2026-10-03")
        self.assertEqual(ps["state_code"], "onway")
        self.assertEqual(ps["orders_cnt"], 8)
        self.assertEqual(ps["dest_name"], "Москва")

        # ② 商品解析（article = vendorCode）
        po = fd.parse_order({
            "id": 5932144473, "nmId": 1357061934, "article": "BCS-OTRE-191868961",
            "title": "Шапка оверсайз", "size": "60-62", "price": "607 ₽",
            "convertedPrice": "46.61 ¥", "stickerID": 58617520011,
            "state": {"code": "shipment", "title": "待发货"},
            "createdAt": "2026-10-02T05:41:25Z", "isCanceled": False, "isTransferred": True,
        }, "WB-GI-287405573")
        self.assertEqual(po["vc"], "BCS-OTRE-191868961")
        self.assertEqual(po["nm_id"], "1357061934")
        self.assertEqual(po["state_code"], "shipment")
        self.assertEqual(po["supply_id"], "WB-GI-287405573")
        self.assertEqual(po["created_day"], "2026-10-02")
        self.assertTrue(po["is_transferred"])
        self.assertFalse(po["is_canceled"])
        # 缺字段容错
        empty = fd.parse_order({}, "S0")
        self.assertEqual((empty["vc"], empty["state_code"], empty["order_id"]), ("", "", ""))

        # ③ 状态过滤：默认只留 shipment；include_canceled 并入 cancel；all = 不过滤
        self.assertEqual(fd.parse_states(""), {"shipment"})
        self.assertEqual(fd.parse_states("shipment,cancel"), {"shipment", "cancel"})
        self.assertEqual(fd.parse_states("all"), set())
        self.assertEqual(fd.parse_states("shipment", include_canceled=True), {"shipment", "cancel"})
        self.assertTrue(fd.match_state("shipment", {"shipment"}))
        self.assertFalse(fd.match_state("sold", {"shipment"}))
        self.assertTrue(fd.match_state("sold", set()))          # 空集 = 不过滤
        # ★ camelCase 状态码必须能被小写的过滤集合命中（实测踩坑：expectedSorting 曾 0 命中）
        self.assertTrue(fd.match_state("expectedSorting", fd.parse_states("expectedSorting")))
        self.assertTrue(fd.match_state("readyForPickup", fd.parse_states("shipment,readyForPickup")))

        # ④ 时间窗口：全空 = 不过滤；闭区间；只给 end 报错
        self.assertEqual(fd.resolve_window(), ("", ""))
        self.assertEqual(fd.resolve_window(date="2026-10-3"), ("2026-10-03", "2026-10-03"))
        self.assertEqual(fd.resolve_window(begin="2026-10-01", end="2026-10-05"),
                         ("2026-10-01", "2026-10-05"))
        with self.assertRaises(ValueError):
            fd.resolve_window(end="2026-10-05")
        self.assertTrue(fd.in_window("2026-10-03", "", ""))     # 不过滤
        self.assertTrue(fd.in_window("2026-10-03", "2026-10-01", "2026-10-05"))
        self.assertFalse(fd.in_window("2026-09-30", "2026-10-01", "2026-10-05"))
        self.assertFalse(fd.in_window("", "2026-10-01", "2026-10-05"))   # 无日期不命中

        # ⑤ 分页：游标推进到底 + 游标不变即停（防死循环）——mock 掉 wb_client.request
        seq = []

        def fake_request(session, method, url, **kw):
            if "supplies/delivery" in url:
                seq.append(kw.get("params", {}).get("next"))
                if len(seq) == 1:
                    return {"data": {"supplies": [{"supplyID": "S1"}], "next": "CUR1"}}
                return {"data": {"supplies": [{"supplyID": "S2"}], "next": 0}}
            return {"data": {"orders": [{"id": 1}], "next": 0}}

        orig = sup.wb_api.request
        try:
            sup.wb_api.request = fake_request
            got = sup.iter_delivery_supplies(object(), max_pages=5)
            self.assertEqual([s["supplyID"] for s in got], ["S1", "S2"])
            self.assertEqual(seq, [0, "CUR1"])                  # 首请求 next=0（对齐抓包），次页用游标

            # 游标恒不变 → 第 2 页后立即停止（不会翻满 max_pages）
            seq.clear()
            sup.wb_api.request = lambda session, method, url, **kw: {
                "data": {"supplies": [{"supplyID": "X"}], "next": "SAME"}}
            got2 = sup.iter_delivery_supplies(object(), max_pages=9)
            self.assertEqual(len(got2), 2)
        finally:
            sup.wb_api.request = orig

        # ⑥ 中文名聚合：商品数降序 + 未归类桶置尾 + vc/清单去重口径
        orders = [
            {"vc": "BCS-AAAA-1", "cn": "甲", "supply_id": "S1", "shop_name": "袁州1", "nm_id": "1"},
            {"vc": "BCS-AAAA-1", "cn": "甲", "supply_id": "S1", "shop_name": "袁州2", "nm_id": "1"},
            {"vc": "BCS-BBBB-2", "cn": "乙", "supply_id": "S2", "shop_name": "袁州1", "nm_id": "2"},
            {"vc": "BCS-CCCC-3", "cn": fd.NO_CN, "supply_id": "S2", "shop_name": "袁州1", "nm_id": "3"},
        ]
        agg = fd.aggregate_by_cn(orders)
        self.assertEqual([(r["cn"], r["item_count"], r["vc_count"], r["supply_count"]) for r in agg],
                         [("甲", 2, 1, 1), ("乙", 1, 1, 1), (fd.NO_CN, 1, 1, 1)])
        self.assertEqual([r["rank"] for r in agg], [1, 2, 3])
        self.assertEqual(agg[0]["shops"], ["袁州1", "袁州2"])
        self.assertTrue(str(agg[-1]["cn"]).startswith("("))     # 未归类桶恒排最后

        # ⑦ 店铺过滤：ID / 短名双认；空 = 全部
        shops = [{"shopId": 9352, "shopName": "袁州1"}, {"shopId": 9353, "shopName": "袁州2"}]
        self.assertEqual([s["shopName"] for s in fd.select_shops(shops, "9352")], ["袁州1"])
        self.assertEqual([s["shopName"] for s in fd.select_shops(shops, "袁州2")], ["袁州2"])
        self.assertEqual(len(fd.select_shops(shops, "9352,袁州2")), 2)
        self.assertEqual(len(fd.select_shops(shops, "")), 2)

    # ---------------- 10. AI 售后客服（买家聊天） ----------------
    def test_47_chats_list_readonly(self):
        """chats 只读命令：列出会话（不发送、不标记已读）。"""
        res = self._run_cmd(["chats", "--shops", "9352", "--limit", "3", "--no-detail"], expect_code=0)
        self.assertIn("会话", res.stdout)

    def test_48_chat_watch_front_once(self):
        """chat-watch 前台模式 --once：只展示不发送（对齐 questions-watch）。"""
        res = self._run_cmd(["chat-watch", "--once", "--shops", "9352"], expect_code=0)
        self.assertIn("--once", res.stdout)

    def test_48b_chatwatch_payload_offline(self):
        """chat-watch 离线断言（**不联网**）：发送体构造 + 上下文轮次整理 + user prompt 组装。

        发送体关键字段（HAR 实测）：sign=会话级 sellerSign、sender.name=店铺名、
        recipients[0].name=买家名、addTime=clientCreatedAt=当前毫秒。
        """
        from wb_ops.adapters import chat_client as cc
        from wb_ops.services.support import chat_prompts as cp

        captured = {}
        orig = cc.wb_api.request

        def _fake(session, method, url, **kw):
            captured.update(method=method, url=url, json=kw.get("json"))
            return {"chatID": "C1", "sign": "SIGN"}

        try:
            cc.wb_api.request = _fake
            d = cc.send_message(object(), "C1", "текст", "Магазин", "Ника", "SIGN", now_ms=123)
        finally:
            cc.wb_api.request = orig

        self.assertEqual(captured["method"], "POST")
        self.assertTrue(captured["url"].endswith("/chat/v5/seller/send"))
        body = captured["json"]
        self.assertEqual(body["chatID"], "C1")
        self.assertEqual(body["sign"], "SIGN")
        self.assertEqual(body["addTime"], 123)
        self.assertEqual(body["clientCreatedAt"], 123)
        self.assertEqual(body["sender"], {"name": "Магазин", "type": "seller"})
        self.assertEqual(body["recipients"], [{"type": "client", "name": "Ника"}])
        self.assertEqual(body["message"]["text"], "текст")
        self.assertEqual(body["message"]["actionType"], "create")
        self.assertEqual(d.get("chatID"), "C1")

        # 会话列表响应顶层即数组 → list_chats 直接返回列表
        cc.wb_api.request = lambda session, method, url, **kw: [{"chatID": "C1"}]
        try:
            self.assertEqual(cc.list_chats(object()), [{"chatID": "C1"}])
        finally:
            cc.wb_api.request = orig

        # recent_turns：倒序 events → 旧→新，排除最新一条
        events = [
            {"isOwn": False, "message": {"text": "m3"}},   # 最新，应排除
            {"isOwn": True, "message": {"text": "m2"}},
            {"isOwn": False, "message": {"text": "m1"}},
        ]
        self.assertEqual(cp.recent_turns(events, max_turns=2), ["买家：m1", "卖家：m2"])
        self.assertEqual(cc.latest_event(events)["message"]["text"], "m3")
        self.assertEqual(cc.latest_seller_sign([{"sellerSign": "SS"}]), "SS")
        self.assertTrue(cc.chat_needs_reply({"lastEvent": {"isOwn": False}}))
        self.assertFalse(cc.chat_needs_reply({"lastEvent": {"isOwn": True}}))

        up = cp.build_user_prompt("商品信息X", "最新问题", cp.recent_turns(events, 2))
        self.assertIn("【需要你回复的买家最新消息】", up)
        self.assertIn("最新问题", up)
        self.assertIn("买家：m1", up)
        # 售后提示词必须含拒退/不暴露 AI 约束
        self.assertIn("绝不", cp.CHAT_SYSTEM_PROMPT)

    def test_55_ratings_readonly(self):
        """测试 ratings 命令只读帮助与核心解析过滤逻辑"""
        res = self._run_cmd(["ratings", "--help"], expect_code=0)
        self.assertIn("ratings", res.stdout)
        self.assertIn("--rating", res.stdout)
        self.assertIn("--one-star", res.stdout)

        from unittest import mock
        from types import SimpleNamespace
        from wb_ops.services.support import ratings

        fake_prods = [
            {
                "nmId": 1334298573,
                "vendorCode": "BCS-PUSQ-90960182",
                "name": "测试商品1",
                "rating": 10,
                "feedbackRating": {"current": 1.0},
                "feedbackCount": {"current": 2},
                "oneStar": {"current": 2},
                "twoStar": {"current": 0},
                "threeStar": {"current": 0},
                "fourStar": {"current": 0},
                "fiveStar": {"current": 0},
            },
            {
                "nmId": 9999999999,
                "vendorCode": "BCS-NONE-11111111",
                "name": "测试商品2",
                "rating": 9,
                "feedbackRating": {"current": 5.0},
                "feedbackCount": {"current": 10},
                "oneStar": {"current": 0},
                "twoStar": {"current": 0},
                "threeStar": {"current": 0},
                "fourStar": {"current": 0},
                "fiveStar": {"current": 10},
            }
        ]

        with mock.patch("wb_ops.adapters.wb_feedback_client.iter_feedback_products", return_value=iter(fake_prods)), \
                mock.patch("wb_ops.services.support.ratings._export_csv") as mock_csv, \
                mock.patch("wb_ops.services.support.ratings._export_excel") as mock_excel:
            args = SimpleNamespace(
                shops="9352",
                rating=1.0,
                min_rating=None,
                max_rating=None,
                one_star=True,
                min_one_star=0,
                min_feedbacks=1,
                days=30,
                name="",
                prefix="",
                vc="",
                limit=0,
                out="",
            )
            ret = ratings.run(args)
            self.assertEqual(ret, 0)
            self.assertTrue(mock_csv.called)
            items = mock_csv.call_args[0][0]
            self.assertEqual(len(items), 1)
            self.assertEqual(items[0].vendor_code, "BCS-PUSQ-90960182")
            self.assertEqual(items[0].feedback_rating, 1.0)
            self.assertEqual(items[0].one_star, 2)
            self.assertEqual(items[0].source_wb, "商品评分API")


    # ---------------- 60. order-stats：订单动态按中文名统计（全离线） ----------------
    def test_60_order_stats_offline(self):
        """order-stats 离线断言：解析/状态归类/窗口/聚合/百分比边界/游标翻页/Sheet 命名/落盘。

        样本形态取自 api/网络请求/wb订单动态 查询退货率.har（2026-10-27，袁州店）：
        状态分布 cancel/buyout/created/return（另有 returnDefective）；
        cursor = base64("<offset>:<时间>")，三页 50/100/150；
        ⚠ currentPeriod 实测按 **updated** 过滤（3 天窗口 65 行中仅 7 行 created 在窗口内）。
        """
        import tempfile
        from datetime import date, datetime
        from unittest import mock

        from wb_ops.adapters import wb_order_feed_client as feed
        from wb_ops.services.order import order_stats as ostats
        from wb_ops.services.order import order_stats_xlsx as oxlsx

        # ① parse_row：完整样本 + 空对象容错
        row = ostats.parse_row({
            "product": {"id": 1585363941, "name": "Шапка зимняя", "subjectName": "保暖帽",
                        "vendorCode": "BCS-ZZMX-1534011939", "productRating": 10},
            "order": {"id": "eBu.r060644b97e4c4173952bc732e9922ce3.0.0",
                      "created": "2026-10-21T01:27:34-07:00",
                      "updated": "2026-10-26T21:20:26-07:00", "status": "buyout",
                      "cancelType": "receipt",
                      "price": {"seller": 49.19}, "size": {"title": "стандарт"},
                      "warehouse": {"title": "склад продавца"},
                      "destination": {"title": "Новосибирск"}},
        })
        self.assertEqual(row["created_day"], "2026-10-21")     # 前缀切片即该时区日期
        self.assertEqual(row["order_no"], "eBu.r060644b97e4c4173952bc732e9922ce3")
        self.assertEqual(row["vc"], "BCS-ZZMX-1534011939")
        self.assertEqual(row["nm_id"], "1585363941")
        self.assertEqual(row["kind"], "done")
        self.assertEqual(row["cancel_type"], "receipt")
        self.assertEqual(row["price"], 49.19)
        empty = ostats.parse_row({})
        self.assertEqual((empty["vc"], empty["status"], empty["created_day"]), ("", "", ""))
        self.assertEqual(ostats.parse_row({"order": {"id": "x"}})["kind"], "unknown")

        # ② classify_status：未知值不猜测
        self.assertEqual(ostats.classify_status("buyout"), "done")
        self.assertEqual(ostats.classify_status("created"), "pending")
        for s in ("cancel", "return", "returnDefective", "RETURN", "Cancel"):
            self.assertEqual(ostats.classify_status(s), "return", s)
        for s in ("delivered", "", None, "sold"):
            self.assertEqual(ostats.classify_status(s), "unknown", s)

        # ③ resolve_window：默认近 30 天含今天（与 fbs-delivery 的"默认不过滤"不同）
        t = date(2026, 10, 27)
        self.assertEqual(ostats.resolve_window(today=t), ("2026-09-28", "2026-10-27"))
        self.assertEqual(ostats.resolve_window(days=1, today=t), ("2026-10-27", "2026-10-27"))
        self.assertEqual(ostats.resolve_window(begin="2026-10-01", end="2026-10-08"),
                         ("2026-10-01", "2026-10-08"))
        self.assertEqual(ostats.resolve_window(date="2026-10-3"), ("2026-10-03", "2026-10-03"))
        self.assertEqual(ostats.resolve_window(begin="2026-10-08", end="2026-10-01"),
                         ("2026-10-01", "2026-10-08"))
        with self.assertRaises(ValueError):
            ostats.resolve_window(end="2026-10-08")
        with self.assertRaises(ValueError):
            ostats.resolve_window(days=1, timezone="Mars/Base")
        # 时区换算（本机无 tzdata，固定偏移 + LA 夏令时）：UTC 2026-10-27 06:00 → LA 10-26
        self.assertEqual(ostats.tz_today("America/Los_Angeles", datetime(2026, 10, 27, 6, 0)),
                         date(2026, 10, 26))
        self.assertEqual(ostats.tz_today("Asia/Shanghai", datetime(2026, 10, 27, 6, 0)),
                         date(2026, 10, 27))

        # ④ aggregate_by_cn：计数/率/未归类桶置尾/分母≤0 → None
        rows = [
            {"cn": "感应灯", "vc": "A", "shop_label": "袁州1(9352)", "kind": "return"},
            {"cn": "感应灯", "vc": "B", "shop_label": "袁州2(9353)", "kind": "done"},
            {"cn": "感应灯", "vc": "B", "shop_label": "袁州2(9353)", "kind": "pending"},
            {"cn": "感应灯", "vc": "B", "shop_label": "袁州2(9353)", "kind": "unknown"},
            {"cn": "", "vc": "C", "shop_label": "袁州1(9352)", "kind": "done"},
            {"cn": "全是未完结", "vc": "D", "shop_label": "袁州1(9352)", "kind": "pending"},
        ]
        agg = ostats.aggregate_by_cn(rows, {"感应灯": "gyd"})
        self.assertEqual([x["cn"] for x in agg],
                         ["感应灯", "全是未完结", "(中文名未知)"])
        top = agg[0]
        self.assertEqual((top["orders"], top["returns"], top["pending"], top["unknown"]), (4, 1, 1, 1))
        self.assertEqual(top["vc_count"], 2)
        self.assertEqual(top["shops"], ["袁州1(9352)", "袁州2(9353)"])
        self.assertEqual(top["stock_sku"], "gyd")
        self.assertAlmostEqual(top["rate"], 0.25)
        self.assertAlmostEqual(top["rate_done"], 1 / 3)
        self.assertEqual([x["rank"] for x in agg], [1, 2, 3])
        pending_only = [x for x in agg if x["cn"] == "全是未完结"][0]
        self.assertIsNone(pending_only["rate_done"])   # 分母=0 → xlsx 写 "-"

        # ④b aggregate_by_vc：跨店同 vc（不同 nmId）必须合并；无 vc 桶置尾；分母≤0 → None
        from wb_ops.services.order import order_stats_vc as ovc
        vc_rows = [
            {"vc": "B", "nm_id": "1001", "cn": "感应灯", "shop_label": "袁州1(9352)", "kind": "return"},
            {"vc": "B", "nm_id": "2002", "cn": "感应灯", "shop_label": "袁州2(9353)", "kind": "done"},
            {"vc": "B", "nm_id": "2002", "cn": "感应灯", "shop_label": "袁州2(9353)", "kind": "pending"},
            {"vc": "A", "nm_id": "9", "cn": "感应灯", "shop_label": "袁州1(9352)", "kind": "done"},
            {"vc": "", "nm_id": "7", "cn": "", "shop_label": "袁州1(9352)", "kind": "done"},
            {"vc": "D", "nm_id": "8", "cn": "全是未完结", "shop_label": "袁州1(9352)", "kind": "pending"},
        ]
        agg_v = ovc.aggregate_by_vc(vc_rows, {"感应灯": "gyd"})
        self.assertEqual(ovc.NO_VC, "(无供应商代码)")
        self.assertEqual([x["vc"] for x in agg_v], ["B", "A", "D", ovc.NO_VC])   # 无 vc 桶恒置尾
        v0 = agg_v[0]
        self.assertEqual((v0["orders"], v0["returns"], v0["pending"], v0["unknown"]), (3, 1, 1, 0))
        self.assertEqual(v0["nm_ids"], ["1001", "2002"])          # 同 vc 跨店两个不同 nmId
        self.assertEqual(v0["shops"], ["袁州1(9352)", "袁州2(9353)"])
        self.assertEqual((v0["cn"], v0["stock_sku"]), ("感应灯", "gyd"))
        self.assertAlmostEqual(v0["rate"], 1 / 3)
        self.assertAlmostEqual(v0["rate_done"], 0.5)
        self.assertEqual([x["rank"] for x in agg_v], [1, 2, 3, 4])
        self.assertEqual(agg_v[3]["cn"], "")                      # --no-cn 场景不伪造 (中文名未知)
        self.assertEqual(agg_v[3]["nm_ids"], ["7"])
        self.assertIsNone([x for x in agg_v if x["vc"] == "D"][0]["rate_done"])   # 分母=0 → None
        # nmId 数值感知排序（纯字典序会把 "999" 排在 "1001" 之后）
        self.assertEqual(ovc.aggregate_by_vc([{"vc": "C", "nm_id": "1001"},
                                              {"vc": "C", "nm_id": "999"}])[0]["nm_ids"], ["999", "1001"])
        # 只给 status 不给 kind ⇒ 走 classify_status（函数内 import，同时验证无循环依赖）
        self.assertEqual(ovc.aggregate_by_vc([{"vc": "E", "status": "cancel"}])[0]["returns"], 1)

        # ⑤ 游标翻页：mock wb_api.request（首页不带 cursor、次页带上一页 cursor）
        bodies = []
        pages = [
            {"data": {"pageInfo": {"cursor": "NTA:MjAyNi0xMC0yNyAwNzoyNToxNg"},
                      "orders": [{"order": {"id": "a.0.0"}}]}},
            {"data": {"pageInfo": {"cursor": ""}, "orders": [{"order": {"id": "b.0.0"}}]}},
        ]

        def fake(session, method, url, **kw):
            bodies.append(kw.get("json"))
            return pages[len(bodies) - 1]

        with mock.patch.object(feed.wb_api, "request", side_effect=fake):
            got = list(feed.iter_orders(object(), "2026-09-28", "2026-10-27", limit=1, sleep=0))
        self.assertEqual([x["order"]["id"] for x in got], ["a.0.0", "b.0.0"])
        self.assertNotIn("cursor", bodies[0])
        self.assertEqual(bodies[1]["cursor"], "NTA:MjAyNi0xMC0yNyAwNzoyNToxNg")
        self.assertEqual(bodies[0]["currentPeriod"], {"start": "2026-09-28", "end": "2026-10-27"})
        self.assertEqual(bodies[0]["timezone"], "America/Los_Angeles")
        self.assertEqual(bodies[0]["orderBy"], {"field": "order.updated", "mode": "desc"})
        self.assertFalse(bodies[0]["isDefective"])
        self.assertEqual(feed.cursor_offset("NTA:MjAyNi0xMC0yNyAwNzoyNToxNg"), 50)

        # cursor 恒不变 → 必须停（防死循环）
        with mock.patch.object(feed.wb_api, "request",
                               return_value={"data": {"pageInfo": {"cursor": "SAME"},
                                                      "orders": [{"order": {"id": "x"}}]}}):
            self.assertEqual(len(list(feed.iter_orders(object(), "a", "b", max_pages=9, sleep=0))), 2)
        # 空页 → 停
        with mock.patch.object(feed.wb_api, "request",
                               return_value={"data": {"pageInfo": {"cursor": "C"}, "orders": []}}):
            self.assertEqual(list(feed.iter_orders(object(), "a", "b", sleep=0)), [])
        # 200 + error=true → 必须抛异常（不能当空数据处理）
        with mock.patch.object(feed.wb_api, "request",
                               return_value={"error": True, "errorText": "bad period"}):
            with self.assertRaises(Exception):
                list(feed.iter_orders(object(), "a", "b", sleep=0))

        # ⑥ 店铺过滤接线（同域复用 fbs_delivery.select_shops）
        shops = [{"shopId": 9352, "shopName": "袁州1"}, {"shopId": 9353, "shopName": "袁州2"}]
        self.assertEqual([s["shopId"] for s in ostats.select_shops(shops, "9352,袁州2")], [9352, 9353])
        self.assertEqual(len(ostats.select_shops(shops, "")), 2)
        self.assertEqual(ostats.select_shops(shops, "9356"), [])

        # ⑦ Sheet 命名：非法字符/截断/重名/撞预留名
        used = {oxlsx.TITLE_SUMMARY.lower(), oxlsx.TITLE_DETAIL.lower()}
        self.assertEqual(oxlsx.safe_sheet_name("袁州1/测试[1]:*?", set()), "袁州1_测试_1____")
        self.assertLessEqual(len(oxlsx.safe_sheet_name("超长" * 40, set())), 31)
        self.assertEqual(oxlsx.safe_sheet_name("汇总", used), "汇总(2)")
        self.assertEqual(oxlsx.safe_sheet_name("袁州1", used), "袁州1")
        self.assertEqual(oxlsx.safe_sheet_name("袁州1", used), "袁州1(2)")
        self.assertEqual(oxlsx.safe_sheet_name(oxlsx.TITLE_VC, {oxlsx.TITLE_VC.lower()}),
                         f"{oxlsx.TITLE_VC}(2)")

        # ⑧ write_xlsx 冒烟：临时目录落盘 → 回读校验结构与格式
        import openpyxl
        detail_row = {"shop_label": "袁州1(9352)", "shop_id": 9352, "order_id": "a.0.0",
                      "order_no": "a", "created_day": "2026-10-21", "created": "2026-10-21T01:27:34-07:00",
                      "updated": "2026-10-26T21:20:26-07:00", "status": "cancel", "kind": "return",
                      "cancel_type": "receipt", "cn": "感应灯", "vc": "BCS-AAAA-1", "stock_sku": "gyd",
                      "nm_id": "1", "title": "Лампа", "size": "стандарт", "price": 49.19,
                      "warehouse": "склад", "dest": "Москва"}
        res = {
            "begin": "2026-09-28", "end": "2026-10-27", "timezone": "America/Los_Angeles",
            "rows": [detail_row],
            "by_cn": [{"rank": 1, "cn": "感应灯", "orders": 1, "returns": 1, "pending": 0,
                       "unknown": 0, "vc_count": 1, "shops": ["袁州1(9352)"],
                       "stock_sku": "gyd", "rate": 1.0, "rate_done": None}],
            "by_cn_by_shop": {"袁州1(9352)": [{"rank": 1, "cn": "感应灯", "orders": 1, "returns": 1,
                                               "pending": 0, "unknown": 0, "vc_count": 1,
                                               "shops": ["袁州1(9352)"], "stock_sku": "gyd",
                                               "rate": 1.0, "rate_done": None}]},
            "by_vc": [{"rank": 1, "vc": "BCS-AAAA-1", "cn": "感应灯", "orders": 1, "returns": 1,
                       "pending": 0, "unknown": 0, "nm_ids": ["1"], "shops": ["袁州1(9352)"],
                       "stock_sku": "gyd", "rate": 1.0, "rate_done": None}],
            "per_shop": [{"shop": "袁州1", "shop_id": 9352, "label": "袁州1(9352)",
                          "orders": 1, "returns": 1, "pending": 0, "unknown": 0, "error": ""}],
            "by_shop": [{"shop_label": "袁州1(9352)", "orders": 1, "returns": 1}],
            "status_counter": {"cancel": 1}, "unknown_status": {}, "out_of_window": 3,
            "failed_shops": [], "dup_order_ids": [],
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = oxlsx.write_xlsx(res, out=os.path.join(tmp, "t.xlsx"))
            self.assertTrue(os.path.exists(path))
            book = openpyxl.load_workbook(path)
            self.assertEqual(book.sheetnames,
                             ["汇总", "按供应商代码", "袁州1(9352)", "订单明细"])
            # 「按供应商代码」Sheet：跨店合并、WB商品码 全量列出、百分比格式、合计行、冻结/换行
            wsv = book[oxlsx.TITLE_VC]
            self.assertEqual([c.value for c in wsv[1]], oxlsx.VC_HEADERS)
            self.assertEqual(len(oxlsx.VC_WIDTHS), len(oxlsx.VC_HEADERS))
            self.assertEqual((wsv["A2"].value, wsv["B2"].value, wsv["C2"].value),
                             (1, "BCS-AAAA-1", "感应灯"))
            self.assertEqual(wsv["F2"].number_format, "0.00%")     # 退货率：数值 + 百分比格式
            self.assertEqual(wsv["G2"].value, "-")                 # 完单口径分母≤0 → 文本 "-"
            self.assertEqual(wsv["J2"].value, "1")                 # WB商品码 全量列出
            self.assertTrue(wsv["J2"].alignment.wrap_text)
            self.assertEqual((wsv["B3"].value, wsv["D3"].value), ("合计", 1))   # 合计行
            self.assertEqual(wsv.freeze_panes, "D2")
            ws = book["汇总"]
            self.assertEqual([c.value for c in ws[1]], oxlsx.HEADERS)
            self.assertEqual(ws["A2"].value, 1)
            self.assertEqual(ws["E2"].number_format, "0.00%")
            self.assertEqual(ws["F2"].value, "-")            # 分母≤0 → 文本 "-"
            self.assertEqual(ws["B3"].value, "合计")          # 合计行
            self.assertEqual(ws["C3"].value, 1)
            self.assertEqual(book["订单明细"].max_row, 2)
            self.assertEqual(book["订单明细"]["I2"].value, "退货")


if __name__ == "__main__":
    unittest.main(verbosity=2)



